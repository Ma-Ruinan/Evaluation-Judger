from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import patch

from docx import Document

from evaluation_judger.reporting import _facts, _fallback, _validate_narrative, compose_narrative, make_formal_report
import json
from evaluation_judger.rubric import METRICS
from evaluation_judger.charts import build_charts
from evaluation_judger.opencode import JudgeError


class LockedReport(TestCase):
    def test_final_failed_review_resumes_its_issues_before_a_new_audit(self):
        facts = {"dataset": "样例", "task_count": 1, "dimension_count": 1,
                 "subjects": {"a": {"name": "甲", "averages": {"completion": "100", "quality_mean": "4"}}},
                 "dimensions": {"维度一": {"tasks": ["1.1"]}}}
        draft = _fallback(facts)
        def answer(value):
            return "BEGIN_JUDGMENT " + json.dumps(value, ensure_ascii=False) + " END_JUDGMENT", False
        with TemporaryDirectory() as directory:
            root = Path(directory)
            responses = [answer(draft),answer({'ok':False,'issues':['首次意见']}),
                         answer(draft),answer({'ok':False,'issues':['最后意见：最大差距归属错误']})]
            with patch('evaluation_judger.reporting.run',side_effect=responses):
                failed = compose_narrative(facts,{},Path(__file__).resolve().parent.parent,root)
            self.assertEqual(failed['source'],'deterministic_fallback')
            review = json.loads((root/'reports/narrative-review.json').read_text(encoding='utf-8'))
            self.assertIn('facts_hash',review)
            self.assertIn('draft_hash',review)
            with patch('evaluation_judger.reporting.run',side_effect=[answer(draft),answer({'ok':True,'issues':[]})]) as model:
                recovered = compose_narrative(facts,{},Path(__file__).resolve().parent.parent,root)
            self.assertEqual(recovered['source'],'opencode')
            self.assertEqual(model.call_count,2)
            self.assertEqual(model.call_args_list[0].args[2],'repair reviewed report prose')
            self.assertIn('最后意见',model.call_args_list[0].args[1])
            self.assertTrue(list((root/'reports/attempt-history').glob('narrative-review*.json')))

    def test_difference_extrema_use_raw_values_and_preserve_ties(self):
        def scores(value):
            return {**{m:{'final':'100' if m=='completion' else value} for m in METRICS},'quality_mean':value}
        rows = [{'task_id':tid,'subjects':{'a':{'scores':scores(a)},'b':{'scores':scores(b)}}}
                for tid,a,b in [('1.1','3','4.063'),('1.2','3','4.3225'),('1.3','4.3225','3')]]
        dims={'维度一':{'tasks':[r['task_id'] for r in rows],'task_results':rows,
             'subjects':{sid:{'scores':{**{m:'100' if m=='completion' else '4' for m in METRICS},'quality_mean':'4'}} for sid in ('a','b')}}}
        facts=_facts(dims,{'dataset':'.','subjects':[{'id':'a','name':'甲'},{'id':'b','name':'乙'}]})
        extrema=facts['dimensions']['维度一']['task_difference_extrema']['b']['quality_mean']
        self.assertEqual(extrema['largest_absolute_difference'],'1.3225')
        self.assertEqual(extrema['task_ids'],['1.2','1.3'])
        self.assertEqual(extrema['positive_task_ids'],['1.1','1.2'])
        self.assertEqual(extrema['negative_task_ids'],['1.3'])

    def test_dimension_mean_differences_are_locked_before_prose_validation(self):
        from decimal import Decimal
        sets = [
            ({'coverage':'4.944','accuracy':'3.3','format':'5','structure':'5','hallucination':'3.523'},
             {'coverage':'5','accuracy':'4.402','format':'5','structure':'5','hallucination':'4.750'}),
            ({'coverage':'5','accuracy':'4.6','format':'4.5','structure':'5','hallucination':'4.13'},
             {'coverage':'5','accuracy':'4.671','format':'5','structure':'4.5','hallucination':'4.39'}),
        ]
        dimensions = {}
        for index,(a,b) in enumerate(sets,1):
            subjects, rows = {}, {}
            for sid,values in [('a',a),('b',b)]:
                scores = {'completion':'100',**values,'quality_mean':str(sum(Decimal(v) for v in values.values())/5)}
                subjects[sid] = {'scores':scores}
                rows[sid] = {'scores':{**{m:{'final':scores[m]} for m in METRICS},'quality_mean':scores['quality_mean']}}
            dimensions[f'维度{index}'] = {'tasks':[f'{index}.1'],'subjects':subjects,'task_results':[{'task_id':f'{index}.1','subjects':rows}]}
        facts = _facts(dimensions, {'dataset':'.','subjects':[{'id':'a','name':'甲'},{'id':'b','name':'乙'}]})
        first = facts['dimensions']['维度1']['display_differences']['b']
        second = facts['dimensions']['维度2']['display_differences']['b']
        self.assertEqual(first['hallucination'], '1.23')
        self.assertEqual(first['quality_mean'], '0.48')
        self.assertEqual(second['hallucination'], '0.26')
        self.assertEqual(second['quality_mean'], '0.07')
        self.assertNotIn('differences', dimensions['维度1'])
        draft = _fallback(facts)
        draft['comparison'] = '维度差值按未舍入锁定数值计算，合法显示为 1.23、0.48、0.26 和 0.07，而不是另行重新评分。'
        _validate_narrative(draft, facts)
        draft['comparison'] += ' 不应接受没有来源的 9.999。'
        with self.assertRaisesRegex(ValueError, 'numbers absent'):
            _validate_narrative(draft, facts)

    def test_invalid_numbers_are_repaired_before_independent_review(self):
        facts = {"dataset": "样例", "task_count": 1, "dimension_count": 1,
                 "subjects": {"a": {"name": "甲", "averages": {"completion": "100", "quality_mean": "4"}}},
                 "dimensions": {"维度一": {"tasks": ["1.1"]}}}
        valid = _fallback(facts)
        invalid = dict(valid, comparison="此处差值 0.07 未包含在锁定事实中，必须由程序拒绝并纠错，不能作为报告发布。")
        def answer(value):
            return "BEGIN_JUDGMENT " + json.dumps(value, ensure_ascii=False) + " END_JUDGMENT", False
        with TemporaryDirectory() as directory:
            root = Path(directory)
            with patch("evaluation_judger.reporting.run", side_effect=[answer(invalid), answer(valid), answer({"ok": True, "issues": []})]) as model:
                result = compose_narrative(facts, {}, Path(__file__).resolve().parent.parent, root)
                again = compose_narrative(facts, {}, Path(__file__).resolve().parent.parent, root)
            self.assertEqual(result["source"], "opencode")
            self.assertEqual(result, again)
            self.assertEqual(model.call_count, 3)
            self.assertIn("0.07", model.call_args_list[1].args[1])
            self.assertEqual(model.call_args_list[2].args[2], "audit locked evaluation report")
            rejected = json.loads((root / "reports/attempt-history/writer-rejected-draft.json").read_text(encoding="utf-8"))
            self.assertEqual(rejected["draft"], invalid)
            self.assertIn("numbers absent", rejected["validation_error"])

    def test_numeric_repair_is_bounded_and_resumes_the_bound_draft(self):
        facts = {"dataset": "样例", "task_count": 1, "dimension_count": 1,
                 "subjects": {"a": {"name": "甲", "averages": {"completion": "100", "quality_mean": "4"}}},
                 "dimensions": {"维度一": {"tasks": ["1.1"]}}}
        valid = _fallback(facts)
        invalid = dict(valid, comparison="此处虚构了不存在于事实表的 9.999 分，应拒绝；有限纠错失败必须保留草稿并停止该阶段。")
        def answer(value):
            return "BEGIN_JUDGMENT " + json.dumps(value, ensure_ascii=False) + " END_JUDGMENT", False
        with TemporaryDirectory() as directory:
            root = Path(directory)
            with patch("evaluation_judger.reporting.run", side_effect=[answer(invalid), answer(invalid)]) as model:
                failed = compose_narrative(facts, {}, Path(__file__).resolve().parent.parent, root)
            self.assertEqual(failed["source"], "deterministic_fallback")
            self.assertEqual(model.call_count, 2)
            self.assertEqual(len(list((root / "reports/attempt-history").glob('writer-rejected-draft*.json'))), 2)
            with patch("evaluation_judger.reporting.run", side_effect=[answer(valid), answer({"ok": True, "issues": []})]) as model:
                recovered = compose_narrative(facts, {}, Path(__file__).resolve().parent.parent, root)
            self.assertEqual(recovered["source"], "opencode")
            self.assertEqual(model.call_count, 2)
            self.assertTrue(model.call_args_list[0].args[2].startswith('repair report validation'))

    def test_report_retries_one_silent_startup_and_reuses_saved_draft(self):
        facts = {"dataset": "样例", "task_count": 1, "dimension_count": 1,
                 "subjects": {"a": {"name": "甲", "averages": {"completion": "100", "quality_mean": "4"}}},
                 "dimensions": {"维度一": {"tasks": ["1.1"]}}}
        draft = _fallback(facts)
        with TemporaryDirectory() as directory:
            root = Path(directory)
            calls = []
            def model(workspace, prompt, title, output, *args, **kwargs):
                calls.append((title, output))
                output.write_text("", encoding="utf-8")
                if len(calls) in (2, 3):
                    raise JudgeError("OpenCode timed out before first event after 180s")
                value = draft if len(calls) == 1 else {"ok": True, "issues": []}
                return "BEGIN_JUDGMENT " + json.dumps(value, ensure_ascii=False) + " END_JUDGMENT", False
            with patch("evaluation_judger.reporting.run", side_effect=model):
                failed = compose_narrative(facts, {}, Path(__file__).resolve().parent.parent, root)
                recovered = compose_narrative(facts, {}, Path(__file__).resolve().parent.parent, root)
            self.assertEqual(failed["source"], "deterministic_fallback")
            self.assertEqual(recovered["source"], "opencode")
            self.assertEqual(len(calls), 4)
            self.assertEqual(sum(title == "write locked evaluation report" for title, _ in calls), 1)
            self.assertEqual(len({path for _, path in calls}), 4)

    def test_real_report_rounded_scores_are_not_treated_as_invented_numbers(self):
        facts = {"dataset": "样例", "task_count": 1, "dimension_count": 1, "subjects": {},
                 "dimensions": {"维度一": {"tasks": ["1.1"]}},
                 "scores": ["4.903225806451612903225806452", "4.940149625935162094763092270",
                            "4.962962962962962962962962963", "4.988746848645900487968510045",
                            "4.99567437750176102182778261", "4.996078431372549019607843138"]}
        draft = _fallback(facts)
        draft["executive_summary"] = "以下分值是锁定结果的舍入表达，未重新评分：4.903、4.940、4.963、4.989、4.9957、4.9961。"
        _validate_narrative(draft, facts)
        for unsupported in ("4.904", "9.999"):
            draft["executive_summary"] = f"以下分值并非锁定结果或其合法舍入，应当继续被程序拒绝：{unsupported}。"
            with self.subTest(value=unsupported), self.assertRaisesRegex(ValueError, "numbers absent"):
                _validate_narrative(draft, facts)

    def test_report_retry_preserves_previous_events_and_archives_resolved_error(self):
        facts = {"dataset": "样例", "task_count": 1, "dimension_count": 1,
                 "subjects": {"a": {"name": "甲", "averages": {"completion": "100", "quality_mean": "4"}}},
                 "dimensions": {"维度一": {"tasks": ["1.1"]}}}
        draft = _fallback(facts)
        def answer(value):
            return "BEGIN_JUDGMENT " + json.dumps(value, ensure_ascii=False) + " END_JUDGMENT", False
        with TemporaryDirectory() as directory:
            root = Path(directory)
            workspace = root / "reports/writer-workspace"
            workspace.mkdir(parents=True)
            original = workspace / "writer-events.jsonl"
            original.write_text("previous attempt", encoding="utf-8")
            error = root / "reports/narrative-error.json"
            error.write_text('{"message":"previous validation error"}', encoding="utf-8")
            with patch("evaluation_judger.reporting.run", side_effect=[answer(draft), answer({"ok": True, "issues": []})]) as model:
                result = compose_narrative(facts, {}, Path(__file__).resolve().parent.parent, root)
            self.assertEqual(result["source"], "opencode")
            self.assertEqual(original.read_text(encoding="utf-8"), "previous attempt")
            self.assertEqual(model.call_args_list[0].args[3].name, "writer-events-1.jsonl")
            self.assertFalse(error.exists())
            archived = root / "reports/attempt-history/narrative-error.json"
            self.assertEqual(json.loads(archived.read_text(encoding="utf-8")), {"message": "previous validation error"})

    def test_review_findings_are_repaired_then_cached_without_rewriting(self):
        facts = {"dataset": "样例数据集", "task_count": 1, "dimension_count": 1,
                 "subjects": {"a": {"name": "甲", "averages": {"completion": "100", "quality_mean": "4"}}},
                 "dimensions": {"维度一": {"tasks": ["1.1"]}}}
        draft = _fallback(facts)
        def answer(value):
            return "BEGIN_JUDGMENT " + json.dumps(value, ensure_ascii=False) + " END_JUDGMENT", False
        with TemporaryDirectory() as directory:
            root = Path(directory)
            responses = [answer(draft), answer({"ok": False, "issues": ["错误类别概括过宽"]}),
                         answer(draft), answer({"ok": True, "issues": []})]
            with patch("evaluation_judger.reporting.run", side_effect=responses) as model:
                first = compose_narrative(facts, {}, Path(__file__).resolve().parent.parent, root)
                second = compose_narrative(facts, {}, Path(__file__).resolve().parent.parent, root)
            self.assertEqual(model.call_count, 4)
            self.assertEqual(first["source"], "opencode")
            self.assertEqual(first, second)
            self.assertTrue((root / "reports/writer-draft.json").exists())

    def test_thirty_task_figures_fit_page_slices_without_losing_numeric_ledger(self):
        rows=[{"task_id":str(i),"subjects":{"a":{"scores":{"quality_mean":"2"}},"b":{"scores":{"quality_mean":"3"}}}} for i in range(30)]
        facts={"task_count":30,"subjects":{sid:{"name":sid,"averages":{m:"2" for m in METRICS}} for sid in ("a","b")},
               "dimensions":{"维度":{"tasks":[str(i) for i in range(30)],"task_results":rows,"subjects":{sid:{"scores":{"quality_mean":"2"}} for sid in ("a","b")}}}}
        with TemporaryDirectory() as directory:
            root=Path(directory);figures=build_charts(facts,root)
            self.assertEqual(len(figures),4)
            self.assertEqual([label for item in figures[2:] for label in item["labels"]],[str(i) for i in range(30)])
            self.assertTrue(all(len(item["labels"])<=15 for item in figures[2:]))
            ledger=json.loads((root/"charts/task-quality-difference.json").read_text(encoding="utf-8"))
            self.assertEqual(ledger["series"]["b"],["1"]*30)

    def test_formal_report_uses_locked_dimension_data(self):
        def scores(value):
            return {**{m: {"final": "100" if m == "completion" else value} for m in METRICS}, "quality_mean": value}

        config = {"dataset": ".", "subjects": [{"id": "a", "name": "甲"}, {"id": "b", "name": "乙"}], "report_title": "试跑报告"}
        dimensions = {"维度一": {
            "tasks": ["1.1"],
            "task_results": [{"task_id": "1.1", "task_name": "样例题", "subjects": {"a": {"scores": scores("4"), "main_losses": []}, "b": {"scores": scores("3"), "main_losses": []}}}],
            "subjects": {sid: {"scores": {**{m: '100' if m=='completion' else value for m in METRICS}, 'quality_mean': value}, "process_stats": {"recorded_success_count": 1, "duration_count": 1, "mean_seconds": "60"}} for sid,value in (("a",'4'), ("b",'3'))},
        }}
        with TemporaryDirectory() as directory:
            with patch("evaluation_judger.reporting.compose_narrative", side_effect=lambda facts, *_: _fallback(facts)):
                target = make_formal_report(config, dimensions, Path(directory), Path(directory))
            doc = Document(target)
            text = "\n".join(p.text for p in doc.paragraphs)
            self.assertIn("执行摘要", text)
            self.assertIn("维度一", text)
            self.assertEqual(_facts(dimensions, config)["subjects"]["a"]["averages"]["quality_mean"], "4")
            self.assertIn("4.00", "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells))
            self.assertNotIn('执行过程表现', text)
            self.assertNotIn('process_stats', json.dumps(_facts(dimensions, config)))
            plot = json.loads((Path(directory)/'reports/charts/task-quality-difference.json').read_text(encoding='utf-8'))
            self.assertEqual(plot['series']['b'], ['-1'])
            self.assertEqual(plot['baseline'], 'a')
            self.assertEqual(len(doc.inline_shapes), 3)
            self.assertIn('附录 逐题六项评分与对象差值', text)
            self.assertTrue(all(p.paragraph_format.keep_with_next for row in doc.tables[-1].rows[1:3] for cell in row.cells for p in cell.paragraphs))
            self.assertFalse(any(p.paragraph_format.keep_with_next for cell in doc.tables[-1].rows[3].cells for p in cell.paragraphs))
