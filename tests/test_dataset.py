from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from evaluation_judger.dataset import discover, process_records, task_fingerprint
from evaluation_judger.judge import workspace_for
import json
from openpyxl import Workbook


def fixture_rubric() -> str:
    lines = ["## 6. 任务完成率"]
    for index, prefix in enumerate("TCAFSH"):
        if index:
            if index == 1:
                lines.append("## 7. 五项质量评分")
            lines.append(f"### 7.{index} test")
        lines += ["| ID | 判据层 | 权重 | 唯一目的 | 状态函数与可观察规则 | 必需证据 | 溯源 |", "| --- | --- | ---: | --- | --- | --- | --- |", f"| {prefix}01 | test | 100 | purpose | **BIN**：rule | file | REQ |"]
    return "\n".join(lines)


class DatasetIsolation(TestCase):
    def test_chinese_dimensions_follow_numeric_or_configured_order(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            for dimension in ("一", "二", "三"):
                for number in (2, 10):
                    task = root / f"维度{dimension}-示例" / f"{number}_题目"
                    task.mkdir(parents=True)
                    (task / "问题描述.txt").write_text("question", encoding="utf-8")
                    (task / "rubric.md").write_text(fixture_rubric(), encoding="utf-8")
                    (task / "a-answer.md").write_text("answer", encoding="utf-8")
            config = {"dataset": str(root), "subjects": [{"id": "a", "name": "A", "prefix": "a-"}]}
            self.assertEqual([task.id for task in discover(config)], ["1.2", "1.10", "2.2", "2.10", "3.2", "3.10"])
            config["dimensions"] = [3, 1]
            self.assertEqual([task.id for task in discover(config)], ["3.2", "3.10", "1.2", "1.10"])

    def test_process_workbook_without_dimension_metadata_is_read_without_modification(self):
        from zipfile import ZipFile
        from xml.etree import ElementTree as ET
        import hashlib
        with TemporaryDirectory() as root:
            path = Path(root) / "process.xlsx"
            wb = Workbook(); ws = wb.active
            ws.append(["task", "time", "status"]); ws.append(["1.1", "2m", "success"])
            wb.save(path)
            with ZipFile(path) as archive:
                parts = {name: archive.read(name) for name in archive.namelist()}
            sheet = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
            dimension = sheet.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}dimension")
            sheet.remove(dimension)
            parts["xl/worksheets/sheet1.xml"] = ET.tostring(sheet)
            with ZipFile(path, "w") as archive:
                for name, data in parts.items(): archive.writestr(name, data)
            before = hashlib.sha256(path.read_bytes()).hexdigest()
            config = {"subjects": [{"id": "a"}], "process_excel": {"path": str(path), "first_row": 2,
                      "last_row": 20, "task_column": "A", "subjects": {"a": {"runtime": "B", "status": "C"}}}}
            self.assertEqual(process_records(config)["1.1"]["a"]["status"], "success")
            config["process_excel"].pop("last_row")
            self.assertEqual(process_records(config)["1.1"]["a"]["runtime"], "2m")
            self.assertEqual(before, hashlib.sha256(path.read_bytes()).hexdigest())

    def test_process_aliases_preserve_original_labels_and_default_numeric_ids(self):
        with TemporaryDirectory() as root:
            path = Path(root) / "process.xlsx"
            wb = Workbook()
            sheet = wb.active
            sheet.append(["task", "time", "status"])
            sheet.append(["W1", "3m", "success"])
            sheet.append(["2.1", "4m", "partial"])
            sheet.append(["1.1", None, None])  # A documentation example, not a test.
            wb.save(path)
            config = {"subjects": [{"id": "a"}], "process_excel": {
                "path": str(path), "task_column": "A", "first_row": 2,
                "task_id_map": {"W1": "1.1"},
                "subjects": {"a": {"runtime": "B", "status": "C"}}}}
            records = process_records(config)
            self.assertEqual(set(records), {"1.1", "2.1"})
            self.assertEqual(records["1.1"]["a"]["task_label"], "W1")
            self.assertEqual(records["2.1"]["a"]["runtime"], "4m")
            sheet.append(["1.1", "5m", "other attempt"])
            wb.save(path)
            with self.assertRaisesRegex(ValueError, "Duplicate process task"):
                process_records(config)

    def test_startup_routing_preserves_fingerprint_and_materials_on_resume(self):
        with TemporaryDirectory() as directory:
            root=Path(directory);folder=root/"data/维度一-演示/1_示例题";folder.mkdir(parents=True)
            (folder/"问题描述.txt").write_text("question",encoding="utf-8")
            (folder/"rubric.md").write_text(fixture_rubric(),encoding="utf-8")
            (folder/"alpha-answer.md").write_text("original answer",encoding="utf-8")
            wb=Workbook();ws=wb.active;ws.append(["1.1","1m","success"]);wb.save(root/"process.xlsx")
            config={"dataset":str(root/"data"),"subjects":[{"id":"alpha","name":"alpha","prefix":"alpha-"}],
                    "process_excel":{"path":str(root/"process.xlsx"),"first_row":1,"task_column":"A","subjects":{"alpha":{"runtime":"B","status":"C"}}}}
            task=discover(config)[0];project=Path(__file__).resolve().parent.parent
            first,original=workspace_for(task,"alpha",config,project,root/"run")
            recovered,manifest=workspace_for(task,"alpha",config,project,root/"run",startup_recovery=True)
            self.assertEqual(manifest["fingerprint"],task_fingerprint(task,"alpha",config))
            self.assertEqual(original,manifest)
            self.assertNotEqual(first,recovered)
            self.assertEqual(workspace_for(task,"alpha",config,project,root/"run")[0],recovered)
            self.assertEqual((recovered/"deliveries/alpha-answer.md").read_text(),"original answer")
            (folder/"alpha-answer.md").write_text("changed answer",encoding="utf-8")
            self.assertNotEqual(workspace_for(task,"alpha",config,project,root/"run")[0],recovered)

    def test_one_to_three_subjects_and_excluded_delivery(self):
        with TemporaryDirectory() as root:
            base = Path(root)
            folder = base / "维度一-演示" / "1_示例题"
            folder.mkdir(parents=True)
            (folder / "问题描述.txt").write_text("question", encoding="utf-8")
            (folder / "rubric.md").write_text(fixture_rubric(), encoding="utf-8")
            (folder / "common.txt").write_text("source", encoding="utf-8")
            for prefix in ("alpha-", "beta-", "gamma-"):
                (folder / f"{prefix}1.1-answer.md").write_text(prefix, encoding="utf-8")
            all_subjects = [{"id": name, "name": name, "prefix": name + "-"} for name in ("alpha", "beta", "gamma")]
            three = discover({"dataset": str(base), "subjects": all_subjects})
            self.assertEqual(len(three), 1)
            self.assertEqual(len(three[0].sources), 1)
            self.assertEqual({k: len(v) for k, v in three[0].deliveries.items()}, {"alpha": 1, "beta": 1, "gamma": 1})
            one = discover({"dataset": str(base), "subjects": all_subjects[:1], "excluded_subject_prefixes": ["beta-", "gamma-"]})
            self.assertEqual([p.name for p in one[0].sources], ["common.txt"])
