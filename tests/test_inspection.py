import json
import hashlib
import os
import shutil
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase, skipUnless

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from evaluation_judger.inspection import allowed_file, inspect, _column
from evaluation_judger.materials import extract


class MaterialInspection(TestCase):
    def test_columns_share_consistent_excel_indexes_and_numeric_header_names(self):
        headers = ["项目", 2024, 2025]
        self.assertEqual(_column(2, headers), 1)
        self.assertEqual(_column("B", headers), 1)
        self.assertEqual(_column("2025", headers), 2)
        for key in (0, -1, 4, "Z"):
            with self.assertRaises(ValueError):
                _column(key, headers)

    @skipUnless(os.name == "nt" and shutil.which("node") and
                (Path(__file__).resolve().parent.parent / ".evaluation-judger/tool-runtime/node_modules/playwright/package.json").exists() and
                Path("C:/Program Files/Google/Chrome/Application/chrome.exe").exists(),
                "Optional pinned Playwright and installed Chrome are required")
    def test_actual_browser_interaction_blocks_external_requests_and_preserves_input(self):
        with TemporaryDirectory() as folder:
            root = Path(folder).resolve()
            file = root / "page.html"
            file.write_text('<!doctype html><meta charset="utf-8"><button id="b" onclick="document.getElementById(\'out\').textContent=\'Changed\'">Switch</button><p id="out">Original</p><script>fetch("https://example.com/private")</script>', encoding="utf-8")
            (root / "manifest.json").write_text(json.dumps({"files": [{"category": "deliveries", "file": "page.html"}]}))
            digest = hashlib.sha256(file.read_bytes()).hexdigest()
            result = inspect(root, {"file": "page.html", "operation": "browser", "widths": [1280, 1920],
                                    "actions": [{"type": "click", "selector": "#b"}]})
            for snapshot in result["snapshots"]:
                self.assertIn("Original", snapshot["initial"]["visible_text"])
                self.assertIn("Changed", snapshot["actions"][0]["after"]["visible_text"])
                self.assertEqual(snapshot["initial"]["document_width"], snapshot["width"])
                self.assertTrue((root / snapshot["screenshot"]).exists())
            self.assertIn("https://example.com/private", result["blocked_requests"])
            self.assertEqual(digest, hashlib.sha256(file.read_bytes()).hexdigest())

    def test_exact_workbook_queries_and_chart_data(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            path = root / "deliveries" / "sample.xlsx"
            path.parent.mkdir()
            wb = Workbook(); ws = wb.active; ws.title = "Data"
            for row in [("kind", "value"), ("x", 1), ("x", 1), ("y", 3)]: ws.append(row)
            ws["C2"] = "=SUM(B2:B4)"
            chart = BarChart(); chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
            ws.add_chart(chart, "E2"); wb.save(path)
            (root / "manifest.json").write_text(json.dumps({"files": [{"category": "deliveries", "file": "deliveries/sample.xlsx"}]}))
            stats = inspect(root, {"file": "deliveries/sample.xlsx", "operation": "stats", "sheet": "Data", "columns": ["value"]})
            self.assertEqual(stats["selected_rows"], 3)
            self.assertEqual(stats["columns"]["value"]["sum"], 5)
            self.assertEqual(stats["columns"]["value"]["duplicate_nonempty"], 1)
            cells = inspect(root, {"file": "deliveries/sample.xlsx", "operation": "cells", "range": "C2:C2"})
            self.assertEqual(cells["cells"][0]["value_or_formula"], "=SUM(B2:B4)")
            package = inspect(root, {"file": "deliveries/sample.xlsx", "operation": "package"})
            self.assertIn("barChart", package["charts"][0]["types"])
            self.assertTrue(any("B" in value for value in package["charts"][0]["references"]))
            filtered = inspect(root, {"file": "deliveries/sample.xlsx", "operation": "rows", "filters": [{"column": "kind", "value": "x"}], "limit": 1})
            self.assertEqual(filtered["selected_rows"], 2)
            self.assertEqual(filtered["rows"][0]["row"], 2)
            self.assertEqual(filtered["next_offset"], 1)

    def test_material_queries_cannot_open_another_participant_or_escape(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "manifest.json").write_text(json.dumps({"files": [{"category": "deliveries", "file": "../other.xlsx"}]}))
            for name in ("opencode.jsonc", "deliveries/competitor.xlsx", "../other.xlsx"):
                with self.assertRaises(ValueError): allowed_file(root, name)

    def test_json_and_long_text_are_paginated_without_discarding_tail(self):
        with TemporaryDirectory() as folder:
            root=Path(folder); file=root/"audit.json"
            file.write_text(json.dumps({"vulns":[{"id":"first"},{"id":"last"}]}))
            (root/"manifest.json").write_text(json.dumps({"files":[{"category":"sources","file":"audit.json"}]}))
            result=inspect(root,{"file":"audit.json","operation":"json","path":["vulns"],"offset":1,"limit":1})
            self.assertEqual(result["value"],[{"id":"last"}]);self.assertIsNone(result["next_offset"])
            text=inspect(root,{"file":"audit.json","operation":"text","offset":10,"limit":10})
            self.assertEqual(text["text"],file.read_text()[10:20])
            with self.assertRaises(ValueError):inspect(root,{"file":"audit.json","operation":"text","offset":-1})

    def test_go_text_and_optional_complete_extraction(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / "version.go"
            path.write_text("package version\n" + "// example\n" * 20000)
            short, limitation = extract(path)
            complete, complete_limitation = extract(path, limit=None)
            self.assertIn("truncated", limitation)
            self.assertIsNone(complete_limitation)
            self.assertGreater(len(complete), len(short))
            self.assertIn("package version", complete)
