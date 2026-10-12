"""Extract readable, locatable text for the model without altering originals."""
from __future__ import annotations

import html
import json
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

TEXT_EXT = {".txt", ".md", ".json", ".jsonl", ".csv", ".tsv", ".py", ".js", ".html", ".htm", ".css", ".xml", ".yaml", ".yml", ".go", ".ts", ".tsx", ".jsx", ".java", ".c", ".h", ".cpp", ".rs", ".sql", ".sh", ".ps1", ".toml"}


def structured_evidence(path: Path) -> dict:
    """Read hidden document structures without executing document content.

    This is an inventory, not a judgment that a source exists or is reliable.
    Empty fields are preserved rather than inferred from neighbouring entries.
    """
    result = {"links": [], "structured_data": [], "limitations": []}
    try:
        if path.suffix.lower() in {".html", ".htm"}:
            soup = BeautifulSoup(path.read_text(encoding="utf-8-sig", errors="replace"), "html.parser")
            for index, anchor in enumerate(soup.find_all("a", href=True), 1):
                result["links"].append({"locator": f"a[href] #{index}",
                    "text": anchor.get_text(" ", strip=True), "target": anchor["href"]})
            for index, script in enumerate(soup.find_all("script"), 1):
                identifier = script.get("id", "")
                if script.get("type", "").lower() not in {"application/json", "application/ld+json"} and not identifier.endswith("-data"):
                    continue
                try:
                    value = json.loads(script.string or script.get_text())
                except (ValueError, TypeError):
                    result["limitations"].append(f"script #{index} {identifier}: not valid JSON; original remains available")
                    continue
                result["structured_data"].append({"locator": f"script #{index} id={identifier}", "data": value})
        elif path.suffix.lower() == ".pdf":
            for page_index, page in enumerate(PdfReader(str(path)).pages, 1):
                for annotation_index, ref in enumerate(page.get("/Annots", []), 1):
                    annotation = ref.get_object()
                    action = annotation.get("/A")
                    if action:
                        action = action.get_object()
                        target = action.get("/URI")
                        if target:
                            result["links"].append({"locator": f"page {page_index}, annotation {annotation_index}",
                                "target": str(target), "rect": [float(x) for x in annotation.get("/Rect", [])]})
    except Exception as exc:
        result["limitations"].append(f"Structure inspection failed: {type(exc).__name__}: {exc}")
    return result


def extract(path: Path, limit: int | None = 160_000) -> tuple[str, str | None]:
    """Return extracted text and an explicit extraction limitation, if any."""
    suffix = path.suffix.lower()
    try:
        if suffix in TEXT_EXT:
            raw = path.read_text(encoding="utf-8-sig", errors="replace")
            if suffix in {".html", ".htm"}:
                soup = BeautifulSoup(raw, "html.parser")
                for element in soup(["script", "style"]):
                    element.decompose()
                raw = html.unescape(soup.get_text("\n", strip=True))
            elif suffix == ".json":
                raw = json.dumps(json.loads(raw), ensure_ascii=False, indent=2)
            text = "\n".join(f"L{i}: {line}" for i, line in enumerate(raw.splitlines(), 1))
        elif suffix == ".pdf":
            reader = PdfReader(str(path))
            text = "\n".join(f"[page {i}]\n{page.extract_text() or ''}" for i, page in enumerate(reader.pages, 1))
        elif suffix == ".docx":
            doc = Document(str(path))
            parts = [f"P{i}: {p.text}" for i, p in enumerate(doc.paragraphs, 1) if p.text.strip()]
            for ti, table in enumerate(doc.tables, 1):
                for ri, row in enumerate(table.rows, 1):
                    parts.append(f"Table{ti} Row{ri}: " + " | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
        elif suffix == ".pptx":
            prs = Presentation(str(path))
            text = "\n".join(f"Slide{i}: {shape.text}" for i, slide in enumerate(prs.slides, 1) for shape in slide.shapes if shape.has_text_frame)
        elif suffix == ".xlsx":
            wb = load_workbook(path, read_only=True, data_only=True)
            parts = []
            for sheet in wb:
                for row in sheet:
                    values = [f"{cell.coordinate}={cell.value}" for cell in row if cell.value is not None]
                    if values:
                        parts.append(f"[{sheet.title}] " + " | ".join(values))
            wb.close()
            text = "\n".join(parts)
        else:
            return "", f"Unsupported file type {suffix}; inspect original file separately"
    except Exception as exc:
        return "", f"Extraction failed: {type(exc).__name__}: {exc}"
    if suffix in {".html", ".htm", ".pdf"}:
        structures = structured_evidence(path)
        if structures["links"] or structures["structured_data"] or structures["limitations"]:
            text += "\n[document links and structured evidence; not source verification]\n" + json.dumps(structures, ensure_ascii=False, indent=2)
    if limit is not None and len(text) > limit:
        return text[:limit], f"Text truncated after {limit} characters; original file remains available"
    return text, None
