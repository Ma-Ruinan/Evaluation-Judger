"""Compose a formal report from locked dimension summaries."""
from __future__ import annotations

import hashlib
import json
import re
import shutil
from decimal import Decimal, ROUND_HALF_UP, localcontext
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .judge import write_json
from .render import _short_reason, _summary_text
from .opencode import JudgeError, extract_json, run
from .rubric import LABELS, METRICS
from .scoring import display, number
from .charts import build_charts
from .isolation import restrict_reads


NARRATIVE_FIELDS = ("executive_summary", "comparison", "limitations", "conclusion")


def _add_prose(doc, text: str):
    """Break long model prose at sentence boundaries for readable Word pages."""
    last = None
    reader_text = str(text).replace('executed_at', '执行时点').replace('ROUND_HALF_UP', '四舍五入')
    for block in reader_text.splitlines():
        pending = ""
        for sentence in re.split(r"(?<=[。！？])", block.strip()):
            if pending and len(pending) + len(sentence) > 350:
                last = doc.add_paragraph(pending)
                pending = ""
            pending += sentence
        if pending:
            last = doc.add_paragraph(pending)
    return last


def _validate_narrative(result: dict, facts: dict) -> None:
    if any(len(str(result.get(key, "")).strip()) < 20 for key in NARRATIVE_FIELDS):
        raise ValueError("Report prose is incomplete")
    analyses = result.get("dimension_analysis")
    if not isinstance(analyses, dict) or any(len(str(analyses.get(dim, "")).strip()) < 20 for dim in facts["dimensions"]):
        raise ValueError("Dimension analysis is incomplete")
    authorized = set(re.findall(r"\d+(?:\.\d+)?", json.dumps(facts, ensure_ascii=False, default=str)))
    numeric_values = {number(value) for value in authorized if len(value.split('.')[0]) < 24}
    authorized.update(display(value) for value in tuple(authorized) if len(value.split('.')[0]) < 24)
    prose = " ".join(str(result[key]) for key in NARRATIVE_FIELDS) + " " + " ".join(str(analyses[dim]) for dim in facts["dimensions"])
    unsupported = set()
    rounded = {}
    with localcontext() as context:
        context.prec = 40
        equivalent = numeric_values | {number(value) for value in authorized if len(value.split('.')[0]) < 24}
        for token in set(re.findall(r"\d+(?:\.\d+)?", prose)) - authorized:
            value = number(token)
            if value in equivalent:
                continue
            precision = len(token.partition('.')[2])
            if 1 <= precision <= 6:
                if precision not in rounded:
                    unit = Decimal(1).scaleb(-precision)
                    rounded[precision] = {item.quantize(unit, rounding=ROUND_HALF_UP) for item in numeric_values}
                if value in rounded[precision]:
                    continue
            unsupported.add(token)
    if unsupported:
        raise ValueError(f"Report prose introduced numbers absent from locked facts: {sorted(unsupported)}")


def _facts(dimensions: dict, config: dict) -> dict:
    dimensions = json.loads(json.dumps(dimensions, ensure_ascii=False, default=str))
    for dimension in dimensions.values():
        for subject in dimension.get('subjects', {}).values():
            subject.pop('process', None)
            subject.pop('process_stats', None)
        for row in dimension.get('task_results', []):
            for subject in row['subjects'].values():
                for entry in subject.get('verification_scope', []):
                    entry['scope_note'] = _summary_text(entry.get('scope_note', ''))
    rows = [row for dimension in dimensions.values() for row in dimension["task_results"]]
    subjects = {}
    for subject in config["subjects"]:
        sid = subject["id"]
        averages = {}
        for metric in METRICS:
            values = [number(row["subjects"][sid]["scores"][metric]["final"]) for row in rows]
            averages[metric] = str(sum(values, Decimal(0)) / len(values))
        averages["quality_mean"] = str(sum((number(averages[m]) for m in METRICS[1:]), Decimal(0)) / 5)
        subjects[sid] = {"name": subject["name"], "averages": averages, "display_averages": {m: display(value) for m, value in averages.items()}}
        if config.get("subject_models", {}).get(sid):
            subjects[sid]["model"] = str(config["subject_models"][sid])
    baseline = config['subjects'][0]['id']
    differences = {s['id']: {m: str(number(subjects[s['id']]['averages'][m])-number(subjects[baseline]['averages'][m])) for m in (*METRICS, 'quality_mean')} for s in config['subjects'][1:]}
    return {
        "dataset": Path(config["dataset"]).name,
        "metric_labels": {**LABELS, "quality_mean": "质量均分"},
        "task_count": len(rows),
        "dimension_count": len(dimensions),
        "subjects": subjects,
        "baseline": baseline,
        "differences": differences,
        "display_differences": {sid: {metric: display(value) for metric,value in values.items()} for sid,values in differences.items()},
        "dimensions": dimensions,
    }


def _fallback(facts: dict) -> dict:
    subject_text = "；".join(f"{item['name']}：完成率 {display(item['averages']['completion'])}%，质量均分 {display(item['averages']['quality_mean'])} 分" for item in facts["subjects"].values())
    return {
        "executive_summary": f"本次纳入 {facts['task_count']} 道题、{facts['dimension_count']} 个维度。{subject_text}。这些数值仅描述本轮纳入的任务。",
        "comparison": "对象差异应结合下方指标表和逐题失分依据阅读；不同题目的核验范围和证据条件可能不同。",
        "limitations": "当前结论只适用于本次纳入的题目和交付。网页可访问性、非文本素材读取范围及历史交付时点会限制部分证据的解释强度。",
        "conclusion": "本报告依据已锁定的逐题判分形成；每项数字可回溯到维度总评和逐题结果。",
        "dimension_analysis": {dim: f"该维度纳入 {len(summary['tasks'])} 道题。主要失分项与对象差异见下方任务表和逐题结果。" for dim, summary in facts["dimensions"].items()},
        "source": "deterministic_fallback",
    }


def _unused_path(directory: Path, stem: str, suffix: str) -> Path:
    """Preserve prior report attempts when retrying the same output stage."""
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{stem}{suffix}"
    index = 1
    while path.exists():
        path = directory / f"{stem}-{index}{suffix}"
        index += 1
    return path


def _report_request(workspace: Path, prompt: str, title: str, stem: str, project: Path, model: str, major: int):
    """Retry one silent startup; preserve events and the bounded stage deadline."""
    for attempt in range(2):
        try:
            return run(workspace, prompt, title, _unused_path(workspace, stem, ".jsonl"), project, model, major=major)
        except JudgeError as exc:
            if attempt or not str(exc).startswith("OpenCode timed out before first event after "):
                raise
            print(f"Retrying report stage after silent startup: {title}", flush=True)


def compose_narrative(facts: dict, config: dict, project: Path, run_dir: Path) -> dict:
    """Use OpenCode for prose; failures leave scores and a usable report intact."""
    report_dir = run_dir / "reports"
    report_dir.mkdir(parents=True, exist_ok=True)
    write_json(report_dir / "report-facts.json", facts)
    digest = hashlib.sha256(json.dumps({"facts": facts, "writer_revision": 9}, ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")).hexdigest()
    checkpoint = report_dir / "narrative.json"
    if checkpoint.exists():
        saved = json.loads(checkpoint.read_text(encoding="utf-8"))
        if saved.get("facts_hash") == digest and saved.get("narrative", {}).get("source") == "opencode":
            return saved["narrative"]
    workspace = report_dir / "writer-workspace"
    workspace.mkdir(exist_ok=True)
    major = int(config.get("opencode_major", 1))
    template = "opencode.v1.example.jsonc" if major == 1 else "opencode.example.jsonc"
    shutil.copy2(project / template, workspace / "opencode.jsonc")
    restrict_reads(workspace)
    write_json(workspace / "locked-facts.json", facts)
    agent_dir = workspace / ".opencode" / "agents"
    agent_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(project / ".opencode" / "agents" / "judge.md", agent_dir / "judge.md")
    prompt = f"""Write the ANALYTIC PROSE for a Chinese AI Agent evaluation report from the locked facts below. Do not score deliveries again. All five quality scores are better when higher; hallucination/self-consistency score measures fewer hallucinations and greater consistency, not hallucination frequency. Follow the Report-Generator principles: executive summary, scope-aware dimension analysis, object comparison, limitations and conclusion. Every concrete difference must cite a dimension/task/atom from the facts. Do not invent task results, numerical values, causes, significance or source claims. A task with no substantive answer may retain nonzero quality scores under explicit empty-set/no-fabrication rubric rules. Such points do not establish completion or useful delivered content; explain completion and quality separately where relevant. For error caps, describe ONLY the concrete finding that triggered the rule: a rule listing several possible error categories does not mean all categories occurred. This may be a subset of the dataset, so state the exact scope. Do not include execution time, runtime status, evaluator operation, success-rate or stability analysis: those records are archived separately and outside this report. Use tables and charts for numbers, prose for specific evidence-backed differences. Avoid internal field names. Keep the executive summary around 150-300 Chinese characters; each dimension analysis around 250-500, other sections around 100-350 as evidence permits. Do not repeat full score vectors in every section: the report tables already show all metrics. Use connected paragraphs in formal Chinese plain prose without Markdown syntax. Return BEGIN_JUDGMENT then a JSON object with keys executive_summary, comparison, limitations, conclusion, dimension_analysis (object keyed by each exact dimension name), then END_JUDGMENT. Facts: read locked-facts.json in THIS workspace using the read tool. It contains the complete locked table as formatted JSON. Read overall scope and scores, then each dimension and its task evidence in segments; continue past any read output cutoff. Do not read outside this workspace or reconstruct judgments from deliveries."""
    prompt = "Distinguish a metric ceiling from the final score: final=min(raw, applicable ceilings). A triggered ceiling does not reduce a raw score already below it. Never call that unchanged score a reduction caused by the cap. An explicit atom-scoped zero cap removes only the listed atoms' contributions and preserves the other atoms in that metric; see scoped_caps if present. " + prompt
    prompt += " Prefer two decimal places for score prose, matching the tables. If finer precision is needed to explain a small difference, use at most six decimal places and ROUND_HALF_UP from the locked value."
    try:
        draft_checkpoint = report_dir / "writer-draft.json"
        saved_draft = json.loads(draft_checkpoint.read_text(encoding="utf-8")) if draft_checkpoint.exists() else {}
        if saved_draft.get("facts_hash") == digest:
            result = saved_draft["draft"]
        else:
            answer, _ = _report_request(workspace, prompt, "write locked evaluation report", "writer-events", project, config.get("model", "aiaaa/deepseek-v4.1-flash#high"), major)
            result = extract_json(answer)
        _validate_narrative(result, facts)
        write_json(draft_checkpoint, {"facts_hash": digest, "draft": result})
        for review_attempt in range(2):
            review_prompt = f"""Act as an independent reviewer of report prose. Compare the draft ONLY against the locked fact table. Check every number, subject mapping, dimension/task/atom attribution, scope limitation and whether any causal or broad superiority claim exceeds the evidence. Do not rescore deliveries. Nonzero quality from explicit empty-set/no-fabrication rules does not prove substantive task completion. Return BEGIN_JUDGMENT then JSON {{"ok":true/false,"issues":[...]}} then END_JUDGMENT. Facts: read locked-facts.json in THIS workspace using the read tool. It contains the complete locked table as formatted JSON. Read overall scope and scores, then each dimension and its task evidence in segments; continue past any read output cutoff. Do not read outside this workspace or reconstruct judgments from deliveries.. Draft: {json.dumps(result, ensure_ascii=False, indent=2)}"""
            review_prompt += " Also explicitly check cap ceilings against raw and final scores: a triggered ceiling above the raw score leaves the score unchanged and cannot be described as causing a reduction."
            review_answer, _ = _report_request(workspace, review_prompt, "audit locked evaluation report", f"review-{review_attempt}-events", project, config.get("model", "aiaaa/deepseek-v4.1-flash#high"), major)
            review = extract_json(review_answer)
            write_json(report_dir / f"narrative-review-{review_attempt}.json", review)
            write_json(report_dir / "narrative-review.json", review)
            if review.get("ok") is True:
                break
            if review_attempt == 1:
                raise ValueError(f"Independent report review found issues: {review.get('issues')}")
            repair_prompt = f"""Revise this Chinese report prose to resolve ONLY the independent review issues. Keep the same JSON keys and all locked scores. Do not infer new errors from a rubric's possible error categories: describe only findings actually recorded. Return BEGIN_JUDGMENT then the complete revised JSON then END_JUDGMENT. Facts: read locked-facts.json in THIS workspace using the read tool. It contains the complete locked table as formatted JSON. Read overall scope and scores, then each dimension and its task evidence in segments; continue past any read output cutoff. Do not read outside this workspace or reconstruct judgments from deliveries.. Draft: {json.dumps(result, ensure_ascii=False, indent=2)}. Issues: {json.dumps(review.get('issues'), ensure_ascii=False, indent=2)}"""
            repaired, _ = _report_request(workspace, repair_prompt, "repair reviewed report prose", "writer-repair-events", project, config.get("model", "aiaaa/deepseek-v4.1-flash#high"), major)
            result = extract_json(repaired)
            _validate_narrative(result, facts)
            write_json(draft_checkpoint, {"facts_hash": digest, "draft": result})
        result["source"] = "opencode"
        prior_error = report_dir / "narrative-error.json"
        if prior_error.exists():
            shutil.copy2(prior_error, _unused_path(report_dir / "attempt-history", "narrative-error", ".json"))
            prior_error.unlink()
    except Exception as exc:
        result = _fallback(facts)
        prior_error = report_dir / "narrative-error.json"
        if prior_error.exists():
            shutil.copy2(prior_error, _unused_path(report_dir / "attempt-history", "narrative-error", ".json"))
        write_json(report_dir / "narrative-error.json", {"type": type(exc).__name__, "message": str(exc)})
    write_json(checkpoint, {"facts_hash": digest, "narrative": result})
    return result


def make_formal_report(config: dict, dimensions: dict, run_dir: Path, project: Path) -> Path:
    facts = _facts(dimensions, config)
    narrative = compose_narrative(facts, config, project, run_dir)
    figures = build_charts(facts, run_dir / 'reports')
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin = sec.bottom_margin = Inches(.72)
    sec.left_margin = sec.right_margin = Inches(.82)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    blue = RGBColor(26, 59, 96)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.name = "Microsoft YaHei"
        doc.styles[name].element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        doc.styles[name].font.color.rgb = blue
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    for border in doc.styles["Title"].element.xpath("./w:pPr/w:pBdr"):
        border.getparent().remove(border)
    cover_title = doc.add_paragraph(config.get("report_title", "AI Agent 应用能力评测报告"), style="Title")
    cover_title.paragraph_format.space_before = Pt(55)
    doc.add_paragraph(f"数据集：{facts['dataset']}  |  纳入 {facts['task_count']} 道题、{facts['dimension_count']} 个维度  |  受测对象 {len(config['subjects'])} 个")
    if config.get('report_scope_note'):
        doc.add_paragraph(str(config['report_scope_note']))
    cover = doc.add_table(rows=1, cols=2)
    cover.style = "Table Grid"
    cover.rows[0].cells[0].text = "测评信息"
    cover.rows[0].cells[1].text = "本轮范围"
    entries = [("数据集", facts["dataset"]), ("评分依据", "各题 Rubric-Council 规则与已锁定的逐题评分结果")]
    entries += [("受测对象 " + str(i+1), item["name"] + (" / " + item["model"] if item.get("model") else "")) for i,item in enumerate(facts["subjects"].values())]
    entries.append(("能力维度", "；".join(facts["dimensions"])))
    for label, value in entries:
        cells = cover.add_row().cells
        cells[0].text, cells[1].text = label, value
    cover.autofit = False
    for column, width in zip(cover.columns, (1.2, 5.35)):
        column.width = Inches(width)
    for row in cover.rows:
        for cell, width in zip(row.cells, (1.2, 5.35)):
            cell.width = Inches(width)
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run(config.get("report_title", "AI Agent 应用能力评测报告")).font.size = Pt(8)
    doc.add_paragraph("报告呈现本轮交付质量、对象差异及可追溯依据，结论限定于上述任务范围。")
    doc.add_page_break()
    doc.add_heading("执行摘要", 1)
    _add_prose(doc, narrative["executive_summary"])
    doc.add_heading("测评范围与方法", 1)
    doc.add_paragraph("依据各题 Rubric-Council 规则独立判分，再从已锁定的逐题结果生成维度总评。质量分由原子权重程序复算；任务等权汇总。下列结论只适用于本报告纳入的题目。")
    models = [f"{item['name']}：{item['model']}" for item in facts["subjects"].values() if item.get("model")]
    if models:
        doc.add_paragraph("本轮登记的受测模型配置：" + "；".join(models) + "。模型名称仅用于说明受测配置，不能据此归因评分差异。")
    doc.add_heading("整体量化结果", 1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, title in zip(table.rows[0].cells, ("受测对象", "任务完成率", "质量均分", "有效题数")):
        cell.text = title
    for item in facts["subjects"].values():
        cells = table.add_row().cells
        for cell, value in zip(cells, (item["name"], display(item["averages"]["completion"]) + "%", display(item["averages"]["quality_mean"]), str(facts["task_count"]))):
            cell.text = value
    doc.add_paragraph("五项质量指标（按本次纳入题目等权平均）")
    metric_table = doc.add_table(rows=1, cols=6)
    metric_table.style = "Table Grid"
    for cell, title in zip(metric_table.rows[0].cells, ("受测对象", *[LABELS[m] for m in METRICS[1:]])):
        cell.text = title
    for item in facts["subjects"].values():
        cells = metric_table.add_row().cells
        for cell, value in zip(cells, (item["name"], *[display(item["averages"][m]) for m in METRICS[1:]])):
            cell.text = value
    precision_note = doc.add_paragraph('统计和差值使用未舍入分数，展示保留两位小数，可能与展示值直接相减略有差别。')
    precision_note.paragraph_format.keep_with_next = True
    doc.add_heading('图表概览', 1)
    for index, figure in enumerate(figures, 1):
        picture = doc.add_paragraph()
        picture.paragraph_format.keep_with_next = True
        picture.add_run().add_picture(figure['path'], width=Inches(6.55))
        caption = doc.add_paragraph(f"图 {index}  {figure['title']}。{figure['caption']}")
        caption.paragraph_format.space_after = Pt(10)
        for run_text in caption.runs:
            run_text.font.size = Pt(9)
    doc.add_heading("各维度分析", 1)
    for dim, summary in dimensions.items():
        doc.add_heading(dim, 2)
        scope = doc.add_paragraph(f"纳入题目：{', '.join(summary['tasks'])}。")
        scope.paragraph_format.keep_with_next = True
        analysis = _add_prose(doc, narrative["dimension_analysis"][dim])
        analysis.paragraph_format.keep_with_next = True
        table = doc.add_table(rows=1, cols=1 + len(config["subjects"]))
        table.style = "Table Grid"
        for cell, value in zip(table.rows[0].cells, ("题目", *[s["name"] + "质量均分" for s in config["subjects"]])):
            cell.text = value
        for row in summary["task_results"]:
            cells = table.add_row().cells
            cells[0].text = f"{row['task_id']} {row['task_name']}"
            for cell, subject in zip(cells[1:], config["subjects"]):
                cell.text = display(row["subjects"][subject["id"]]["scores"]["quality_mean"])
        evidence_rows = []
        for row in summary["task_results"]:
            for subject in config["subjects"]:
                item = row["subjects"][subject["id"]]
                for loss in item.get("main_losses", [])[:1]:
                    evidence_rows.append((number(item["scores"]["quality_mean"]),
                                          f"{row['task_id']}/{loss['atom_id']} · {subject['name']}：{loss['brief']}"))
        if evidence_rows:
            doc.add_paragraph("代表性失分依据").paragraph_format.keep_with_next = True
        for _, text in sorted(evidence_rows, key=lambda entry: entry[0])[:4]:
            doc.add_paragraph(text, style="List Bullet").paragraph_format.keep_together = True
        cap_rows = [(row, subject, error) for row in summary["task_results"] for subject in config["subjects"]
                    for error in row["subjects"][subject["id"]].get("critical_errors", []) if error.get("triggered")]
        for row, subject, error in cap_rows:
            doc.add_paragraph(f"{row['task_id']}/{error['id']} · {subject['name']}：关键错误触发封顶。{_short_reason(error['reason'], 180)}", style="List Bullet").paragraph_format.keep_together = True
    if len(config["subjects"]) > 1:
        doc.add_heading("对象比较与主要差异", 1)
        _add_prose(doc, narrative["comparison"])
    doc.add_heading("限制与适用范围", 1)
    _add_prose(doc, narrative["limitations"])
    doc.add_heading("结论", 1)
    _add_prose(doc, narrative["conclusion"]).paragraph_format.keep_together = True
    appendix = doc.add_section(WD_SECTION_START.NEW_PAGE)
    appendix.orientation = WD_ORIENT.LANDSCAPE
    appendix.page_width, appendix.page_height = Mm(297), Mm(210)
    doc.add_heading('附录 逐题六项评分与对象差值', 1)
    doc.add_paragraph('完成率单位为百分比，完成率差值为百分点；其余为分。差值使用未舍入分数计算。质量均分只包含五项质量指标。')
    if len(config['subjects']) > 1:
        doc.add_paragraph('基准对象：' + config['subjects'][0]['name'] + '。差值为其他对象减基准对象。')
    for dim, summary in dimensions.items():
        doc.add_heading(dim, 2)
        table = doc.add_table(rows=1, cols=9)
        table.style = 'Table Grid'
        table.autofit = False
        widths = [.6, 2.1, 1.05, 1, 1, 1, 1, 1, 1]
        for cell, width, label in zip(table.rows[0].cells, widths, ('题号', '对象或差值', *[LABELS[m] for m in METRICS], '质量均分')):
            cell.width = Inches(width)
            cell.text = label
        for row in summary['task_results']:
            group_start = len(table.rows)
            base = config['subjects'][0]
            for subject in config['subjects']:
                scores = row['subjects'][subject['id']]['scores']
                values = [row['task_id'], subject['name'], *[display(scores[m]['final'])+('%' if m=='completion' else '') for m in METRICS], display(scores['quality_mean'])]
                for cell, width, value in zip(table.add_row().cells, widths, values):
                    cell.width = Inches(width)
                    cell.text = value
            for subject in config['subjects'][1:]:
                scores = row['subjects'][subject['id']]['scores']
                baseline = row['subjects'][base['id']]['scores']
                values = [row['task_id'], subject['name']+'差值', *[display(number(scores[m]['final'])-number(baseline[m]['final'])) for m in METRICS], display(number(scores['quality_mean'])-number(baseline['quality_mean']))]
                for cell, width, value in zip(table.add_row().cells, widths, values):
                    cell.width = Inches(width)
                    cell.text = value
            # Keep one task's subject rows and delta rows on the same page.
            for grouped_row in table.rows[group_start:-1]:
                for cell in grouped_row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = "能力测评报告 · 第 "
    page_number = OxmlElement("w:fldSimple")
    page_number.set(qn("w:instr"), "PAGE")
    footer._p.append(page_number)
    footer.add_run(" 页")
    for table in doc.tables:
        header = table.rows[0]
        repeat = OxmlElement("w:tblHeader")
        header._tr.get_or_add_trPr().append(repeat)
        for cell in header.cells:
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "EAF0F6")
            cell._tc.get_or_add_tcPr().append(shade)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
                for text_run in paragraph.runs:
                    text_run.bold = True
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    target = run_dir / "reports" / "evaluation-report.docx"
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    return target
