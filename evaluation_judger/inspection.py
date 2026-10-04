"""Read-only material queries and controlled document/browser previews."""
from __future__ import annotations

import argparse
from collections import Counter
import json
import math
from pathlib import Path
import statistics
import sys
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, range_boundaries


def allowed_file(workspace: Path, name: str) -> Path:
    root = workspace.resolve()
    manifest = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    allowed = {entry["file"] for entry in manifest["files"]
               if entry["category"] in {"question", "rubric", "sources", "deliveries"}}
    if name not in allowed:
        raise ValueError("File is not an authorized material in this task-participant workspace")
    target = (root / name).resolve()
    if not target.is_relative_to(root):
        raise ValueError("Material path leaves the isolated workspace")
    return target


def package_features(path: Path) -> dict:
    """Expose actual OOXML chart definitions, never infer visual appearance."""
    with ZipFile(path) as archive:
        names = archive.namelist()
        charts = []
        ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
              "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        for name in names:
            if "/charts/chart" not in name or not name.endswith(".xml"):
                continue
            root = ET.fromstring(archive.read(name))
            charts.append({"part": name, "types": sorted({node.tag.split("}")[-1]
                           for node in root.iter() if node.tag.split("}")[-1].endswith("Chart")}),
                           "text": [node.text for node in root.findall(".//a:t", ns) if node.text],
                           "references": [node.text for node in root.findall(".//c:f", ns)],
                           "cached_values": [node.text for node in root.findall(".//c:v", ns)],
                           "definition_xml": archive.read(name).decode("utf-8")})
        return {"charts": charts, "xml_parts": [n for n in names if n.endswith(".xml")],
                "images": [n for n in names if "/media/" in n and not n.endswith("/")],
                "pivot_parts": [n for n in names if "/pivotTables/" in n and n.endswith(".xml")],
                "limitation": "Chart XML proves content and settings, not final rendered appearance."}


def _numeric(value):
    if isinstance(value, bool) or value is None:
        return None
    try:
        number = float(value)
        return number if math.isfinite(number) else None
    except (TypeError, ValueError):
        return None


def _column(key, headers) -> int:
    if isinstance(key, int) and not isinstance(key, bool):
        if not 1 <= key <= len(headers):
            raise ValueError("Integer column indexes are 1-based and must be inside the worksheet")
        return key - 1
    if key in headers:
        return headers.index(key)
    for index, header in enumerate(headers):
        if header is not None and str(header) == str(key):
            return index
    index = column_index_from_string(str(key).upper()) - 1
    if not 0 <= index < len(headers):
        raise ValueError("Column is outside the worksheet header span")
    return index


def inspect(workspace: Path, request: dict) -> dict:
    path = allowed_file(workspace, request["file"])
    operation = request.get("operation", "summary")
    if operation in {"pptx_visual", "browser"}:
        from evaluation_judger.visual import powerpoint, browser
        return (powerpoint if operation == "pptx_visual" else browser)(workspace.resolve(), path, request)
    if operation == "json":
        if path.suffix.lower() != ".json":
            raise ValueError("JSON queries require a .json material")
        selected = json.loads(path.read_text(encoding="utf-8-sig"))
        keys = request.get("path", [])
        if not isinstance(keys, list):
            raise ValueError("JSON path must be an array of keys or integer indices")
        for key in keys:
            selected = selected[key]
        offset, limit = int(request.get("offset", 0)), int(request.get("limit", 25))
        if offset < 0 or not 1 <= limit <= 100:
            raise ValueError("offset >= 0 and 1 <= limit <= 100 required")
        if isinstance(selected, dict):
            values = list(selected.items())
            value = dict(values[offset:offset+limit])
        elif isinstance(selected, list):
            values = selected
            value = selected[offset:offset+limit]
        else:
            return {"path": keys, "value": selected}
        return {"path": keys, "total": len(values), "offset": offset, "value": value,
                "next_offset": offset+limit if offset+limit<len(values) else None}
    if operation == "text":
        if path.suffix.lower() not in {".txt", ".md", ".json", ".csv", ".html", ".css", ".js", ".ts", ".py", ".go", ".sql", ".yaml", ".yml", ".toml", ".xml", ".sh", ".ps1"}:
            raise ValueError("Text queries require a supported text material")
        offset, limit = int(request.get("offset", 0)), int(request.get("limit", 8000))
        if offset < 0 or not 1 <= limit <= 20000:
            raise ValueError("offset >= 0 and 1 <= limit <= 20000 required")
        content = path.read_text(encoding="utf-8-sig")
        return {"offset": offset, "total_characters": len(content), "starting_line": content[:offset].count("\n")+1,
                "text": content[offset:offset+limit], "next_offset": offset+limit if offset+limit<len(content) else None}
    if operation == "xml":
        offset,limit=int(request.get("offset",0)),int(request.get("limit",200))
        if offset<0 or not 1<=limit<=400:
            raise ValueError("offset >= 0 and 1 <= limit <= 400 required")
        with ZipFile(path) as archive:
            part=request["part"]
            if part not in archive.namelist() or not part.endswith(".xml"):
                raise ValueError("Select an existing XML part from the package inventory")
            root=ET.fromstring(archive.read(part));ET.indent(root)
            lines=ET.tostring(root,encoding="unicode").splitlines()
            return {"part":part,"total_lines":len(lines),"lines":[{"line":i+1,"text":s} for i,s in enumerate(lines[offset:offset+limit],offset)],
                    "next_offset":offset+limit if offset+limit<len(lines) else None}
    if operation == "package":
        return package_features(path)
    if path.suffix.lower() != ".xlsx":
        if path.suffix.lower() in {".docx", ".pptx"}:
            return package_features(path)
        raise ValueError("Use the read tool for text; workbook queries require .xlsx")
    if operation == "cells":
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = wb[request.get("sheet", wb.sheetnames[0])]
            bounds = range_boundaries(request.get("range", "A1:J20"))
            if (bounds[2] - bounds[0] + 1) * (bounds[3] - bounds[1] + 1) > 4000:
                raise ValueError("At most 4000 cells per query; paginate ranges")
            return {"sheet": sheet.title, "cells": [{"cell": c.coordinate, "value_or_formula": c.value,
                    "format": c.number_format} for row in sheet.iter_rows(min_col=bounds[0], min_row=bounds[1],
                    max_col=bounds[2], max_row=bounds[3]) for c in row if c.value is not None]}
        finally:
            wb.close()
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if operation == "summary" and "sheet" not in request:
            return {"sheets": [{"name": s.title, "rows": s.max_row, "columns": s.max_column,
                                "state": s.sheet_state} for s in wb], **package_features(path)}
        sheet = wb[request.get("sheet", wb.sheetnames[0])]
        header_row = int(request.get("header_row", 1))
        if header_row < 1:
            raise ValueError("header_row must be positive")
        headers = list(next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True)))
        rows = [(i, tuple(row)) for i, row in enumerate(sheet.iter_rows(min_row=header_row+1, values_only=True), header_row+1)
                if any(value is not None for value in row)]
        for condition in request.get("filters", []):
            index = _column(condition["column"], headers)
            op, operand = condition.get("op", "eq"), condition.get("value")
            def matches(row):
                value = row[index]
                if op == "eq": return value == operand
                if op == "ne": return value != operand
                if op == "empty": return value is None or str(value).strip() == ""
                if op == "not_empty": return value is not None and str(value).strip() != ""
                if op == "contains": return str(operand) in str(value)
                if op in {"lt", "gt"}:
                    number, threshold = _numeric(value), _numeric(operand)
                    return number is not None and threshold is not None and (number < threshold if op == "lt" else number > threshold)
                raise ValueError("Unsupported filter operation")
            rows = [(i, row) for i, row in rows if matches(row)]
        base = {"sheet": sheet.title, "headers": headers, "header_row": header_row,
                "selected_rows": len(rows), "duplicate_rows": len(rows) - len({row for _, row in rows})}
        if operation == "rows":
            offset, limit = int(request.get("offset", 0)), int(request.get("limit", 100))
            if offset < 0 or not 1 <= limit <= 200:
                raise ValueError("offset >= 0 and 1 <= limit <= 200 required")
            return {**base, "rows": [{"row": i, "values": row} for i, row in rows[offset:offset+limit]],
                    "next_offset": offset+limit if offset+limit < len(rows) else None}
        if operation in {"summary", "stats"}:
            columns = request.get("columns", [str(h) if h is not None else i+1 for i,h in enumerate(headers)])
            result = {}
            for column in columns:
                index = _column(column, headers)
                values = [row[index] for _, row in rows]
                present = [v for v in values if v is not None and str(v).strip()]
                numeric = [n for v in present if (n := _numeric(v)) is not None]
                counter = Counter(present)
                result[str(column)] = {"nonempty": len(present), "empty": len(values)-len(present),
                    "distinct": len(counter), "duplicate_nonempty": len(present)-len(counter),
                    "numeric_count": len(numeric), "noninteger_count": sum(not n.is_integer() for n in numeric),
                    "sum": sum(numeric) if numeric else None, "mean": statistics.mean(numeric) if numeric else None,
                    "min": min(numeric) if numeric else None, "max": max(numeric) if numeric else None,
                    "top_values": [{"value": v, "count": n} for v,n in counter.most_common(25)]}
            return {**base, "columns": result}
        if operation == "correlation":
            x,y = (_column(request[key], headers) for key in ("x", "y"))
            pairs = [(a,b) for _,r in rows if (a:=_numeric(r[x])) is not None and (b:=_numeric(r[y])) is not None]
            return {**base, "pair_count": len(pairs), "pearson": statistics.correlation(*zip(*pairs)) if len(pairs)>1 else None}
        if operation == "group":
            index = _column(request["group_by"], headers)
            groups = {}
            for _,row in rows: groups.setdefault(row[index], []).append(row)
            results = []
            for key,group in groups.items():
                entry = {"group": key, "row_count": len(group)}
                for column in request.get("columns", []):
                    ci = _column(column,headers);values=[v for r in group if (v:=_numeric(r[ci])) is not None]
                    entry[str(column)]={"count":len(values),"sum":sum(values),"mean":statistics.mean(values) if values else None}
                results.append(entry)
            return {**base,"groups":results}
        raise ValueError("Unsupported operation")
    finally:
        wb.close()


def install_tool(workspace: Path) -> None:
    # The native MCP path avoids loading an external Zod/plugin module into
    # OpenCode's tool registry. Keep the original inspection implementation.
    root = workspace.resolve()
    config_path = root / "opencode.jsonc"
    config = json.loads(config_path.read_text(encoding="utf-8-sig"))
    config.setdefault("mcp", {})["material"] = {
        "type": "local", "enabled": True, "timeout": 130000,
        "command": [sys.executable, "-X", "utf8", str(Path(__file__).with_name("inspection_mcp.py").resolve()),
                    "--workspace", str(root)],
    }
    for extension in ("ts", "js"):
        wrapper = root / f".opencode/tools/material_inspector.{extension}"
        if wrapper.exists():
            archived = wrapper.with_suffix(f".{extension}.disabled")
            if archived.exists():
                if archived.read_bytes() != wrapper.read_bytes():
                    raise ValueError("Existing disabled inspector wrapper differs; retain it before migration")
                wrapper.unlink()
            else:
                wrapper.rename(archived)
    config_path.write_text(json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8")


def main():
    parser=argparse.ArgumentParser();parser.add_argument("--workspace",type=Path,required=True);parser.add_argument("--request",required=True)
    args=parser.parse_args()
    try:
        print(json.dumps(inspect(args.workspace,json.loads(args.request)),ensure_ascii=False,default=str))
    except Exception as exc:
        print(json.dumps({"error":type(exc).__name__,"message":str(exc)},ensure_ascii=False));raise SystemExit(1)


if __name__ == "__main__":main()
