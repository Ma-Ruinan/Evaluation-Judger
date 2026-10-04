"""Render traceable Markdown, dimension summaries, and a dataset Word report."""
from __future__ import annotations

import json
import hashlib
import re
from html import escape
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from .dataset import process_records
from .judge import write_json, write_text
from .rubric import LABELS, METRICS
from .scoring import display, number


def _safe(value) -> str:
    # Markdown permits raw HTML; an unclosed <script> can hide the rest of a report.
    return escape(str(value or ""), quote=False).replace("|", "\\|").replace("\n", " ")


def _item_label(item: dict) -> str:
    return str(item.get("claim") or item.get("element") or item.get("statement") or item.get("entry") or item.get("name") or item.get("id") or "")


def _item_passed(item: dict) -> bool | None:
    outcome = str(item.get("verdict", item.get("result", ""))).strip().lower()
    if outcome.startswith(("不满足", "不一致", "不支持", "无支撑", "未覆盖", "缺失", "错误", "冲突", "矛盾")):
        return False
    if outcome.startswith(("满足", "一致", "支持", "有支撑", "已覆盖", "正确", "可追溯", "非虚构", "无问题", "无冲突")):
        return True
    if outcome.startswith(("satisfied", "supported", "consistent", "correct", "traceable", "pass", "present", "yes")):
        return True
    if outcome.startswith(("not satisfied", "unsatisfied", "unsupported", "inconsistent", "incorrect", "fail", "missing", "absent", "no")):
        return False
    if outcome in ("1", "true"):
        return True
    if outcome in ("0", "false"):
        return False
    return None


def _visible_items(items: list) -> tuple[list, int]:
    """Keep all negative checks visible; full item ledger remains in JSON."""
    if len(items) <= 10:
        return items, 0
    selected, positives = [], 0
    for item in items:
        status = _item_passed(item) if isinstance(item, dict) else None
        # An unfamiliar label could encode a deduction; omit only positives.
        if status is False or status is None or status is True and positives < 2:
            selected.append(item)
            positives += status is True
    return selected, len(items) - len(selected)


def _atom_difference(atom_records: list[dict], names: list[str]) -> str:
    """Compare only shared, identifiable checks; different samples imply no absence."""
    if len(atom_records) != 2:
        return "分数见上表；完整依据见下文。"
    left, right = atom_records
    right_items = {_item_label(item).strip(): item for item in right.get("items", []) if isinstance(item, dict) and _item_label(item).strip()}
    differences = []
    for item in left.get("items", []):
        if not isinstance(item, dict):
            continue
        label = _item_label(item).strip()
        partner = right_items.get(label)
        if partner is None:
            continue
        a, b = _item_passed(item), _item_passed(partner)
        if a is not None and b is not None and a != b:
            winner, loser = (names[0], names[1]) if a else (names[1], names[0])
            differences.append(f"{_safe(label)}：{_safe(winner)}满足，{_safe(loser)}未满足")
    if differences:
        tail = f"；另有 {len(differences) - 2} 个同项差异见逐项依据" if len(differences) > 2 else ""
        return "；".join(differences[:2]) + tail
    parts = []
    for name, atom in zip(names, atom_records):
        failed = [item for item in atom.get("items", []) if isinstance(item, dict) and _item_passed(item) is False]
        if failed:
            examples = "、".join(_safe(_item_label(item) or item.get("source") or "核验项") for item in failed[:2])
            parts.append(f"{_safe(name)}：{len(failed)}项未通过，如{examples}")
        elif number(atom["state"]) == 1:
            observation = atom.get("observation") or atom.get("reason") or "本项满足规则"
            parts.append(f"{_safe(name)}：{_safe(_short_reason(observation, 145))}")
        else:
            parts.append(f"{_safe(name)}：{_safe(_short_reason(atom.get('reason') or atom.get('observation', ''), 145))}")
    suffix = "。具体核验清单和判定见下文。" if left.get("items") and right.get("items") else "。"
    return "；".join(parts) + suffix


def duration_seconds(value) -> int | None:
    if value is None:
        return None
    raw = str(value).strip().lower()
    match = re.fullmatch(r"(?:(\d+)h)?(?:(\d+)m)?(?:(\d+)s)?", raw)
    if match and any(x is not None for x in match.groups()):
        h, m, s = (int(x or 0) for x in match.groups())
        return h * 3600 + m * 60 + s
    return None


def _summary_text(value: str) -> str:
    """Exclude embedded operational metadata from exported analytical prose."""
    return re.sub(r'process-record\.json[^\n；;]*(?:runtime|status)\s*=[^\n；;]*[；;]?', '', str(value), flags=re.I).strip()


def _short_reason(value: str, limit: int = 145) -> str:
    cleaned = " ".join(str(value).split())
    cleaned = re.sub(r"^(?:BIN|RATIO|CLAIM-RATIO|COUNT)(?:规则|判据)?[：:]\s*", "", cleaned)
    cleaned = re.sub(r"(?:numerator|denominator|state)\s*=\s*[\d.]+[，,、;；]?", "", cleaned, flags=re.I)
    if len(cleaned) <= limit:
        return cleaned
    candidate = cleaned[:limit]
    end = max(candidate.rfind("。"), candidate.rfind("；"))
    return candidate[:end + 1] if end > 65 else candidate.rstrip("，、；。 ") + "……"


def _loss_brief(atom: dict) -> str:
    """Show the actual failed checks in the report, not a clipped rubric recital."""
    items = atom.get("items") or []
    failed = []
    for item in items:
        if not isinstance(item, dict):
            continue
        if _item_passed(item) is False:
            failed.append(item)
    ratio = ""
    if atom.get("denominator") is not None:
        ratio = f"{atom['numerator']}/{atom['denominator']}；"
    if failed:
        examples = []
        for item in failed[:2]:
            name = str(item.get("claim") or item.get("element") or item.get("source") or item.get("id") or "核验项")
            outcome = str(item.get("verdict", item.get("result", ""))).lower().strip()
            note = str(item.get("note") or item.get("reason") or "")
            if not note:
                note = ("页面无法取得，来源及内容未获证实" if outcome.startswith("not satisfied") else
                        "缺少可识别的来源支撑" if outcome.startswith("unsupported") else
                        "与已核来源不一致")
            location = str(item.get("location") or item.get("scope") or "")
            if location.startswith(("extracted/", "deliveries/")):
                location = ""
            examples.append(f"{name}（{location}）：{_short_reason(note, 85)}" if location else f"{name}：{_short_reason(note, 85)}")
        remainder = f"；另有 {len(failed) - 2} 项，详见逐题结果" if len(failed) > 2 else ""
        return f"{ratio}{'；'.join(examples)}{remainder}。"
    # A binary reason may explain a satisfied condition before the failed one.
    # Preserve its full rationale rather than clipping away the actual finding.
    return f"{ratio}{_summary_text(atom['reason'])}"


def task_markdown(task, results: list[dict], names: dict[str, str]) -> str:
    lines = [f"# {task.id} {_safe(task.name)} 评测结果", "", f"维度：{_safe(task.dimension)}", "", "## 分数", ""]
    headers = ["指标", *[_safe(names[r["subject"]]) for r in results]]
    if len(results) > 1:
        headers += [f"{_safe(names[r['subject']])}−{_safe(names[results[0]['subject']])}" for r in results[1:]]
    lines += ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    for metric in METRICS:
        values = [number(r["scores"][metric]["final"]) for r in results]
        cells = [display(v) + ("%" if metric == "completion" else "") for v in values]
        cells += [display(v - values[0]) for v in values[1:]]
        lines.append("| " + LABELS[metric] + " | " + " | ".join(cells) + " |")
    values = [number(r["scores"]["quality_mean"]) for r in results]
    lines.append("| 质量均分 | " + " | ".join([*[display(v) for v in values], *[display(v - values[0]) for v in values[1:]]]) + " |")
    cap_notes = []
    for result in results:
        errors = [error for error in result["errors"] if error.get("triggered")]
        if errors:
            changes = [f"{LABELS[m]} {display(result['scores'][m]['raw'])} → {display(result['scores'][m]['final'])}"
                       for m in METRICS if result["scores"][m]["caps"] or result["scores"][m].get("scoped_caps")]
            changes += [f"{LABELS[m]}仅对 {'、'.join(scope['atoms'])} 应用局部上限 {scope['ceiling']}"
                        for m in METRICS for scope in result['scores'][m].get('scoped_caps', [])]
            cap_notes.append(f"- {_safe(names[result['subject']])}：触发 {'、'.join(error['id'] for error in errors)}；{'；'.join(changes)}。具体依据见该对象的关键错误与封顶记录。")
    if cap_notes:
        lines += ["", "封顶说明：上表为应用关键错误上限后的最终分数。", "", *cap_notes, ""]
    if len(results) > 1:
        lines += ["", "比较口径：各对象分别按同一 rubric 独立判分；差值为后列对象减首列对象，按未舍入分数计算。", ""]
        lines += ["## 逐项差异速览", "", "仅对齐同名核验项；不同抽样清单不据此推断另一对象缺失。", "", "| 打分项 | 状态差异 | 可核对的具体差异 |", "| --- | --- | --- |"]
        for rubric_atom in task.rubric.atoms:
            pair = [next(a for a in result["atoms"] if a["id"] == rubric_atom.id) for result in results]
            states = [display(a["state"]) for a in pair]
            if len(set(states)) == 1:
                continue
            summary = "；".join(_atom_difference([pair[0], pair[index]], [names[results[0]["subject"]], names[results[index]["subject"]]])
                                 for index in range(1, len(results)) if number(pair[0]["state"]) != number(pair[index]["state"]))
            lines.append(f"| {rubric_atom.id}（{LABELS[rubric_atom.metric]}） | {' / '.join(states)} | {summary} |")
        lines.append("")
    for result in results:
        lines += ["", f"## {_safe(names[result['subject']])} 逐项依据", "", f"完整逐项核验清单：`{result['subject']}.json`。下文对较长清单展示全部未通过项及少量通过项示例。", ""]
        for metric in METRICS:
            lines += [f"### {LABELS[metric]}", ""]
            for atom in [a for a in result["atoms"] if a["metric"] == metric]:
                lines += [f"#### {atom['id']} 状态 {display(atom['state'])} 权重 {atom['weight']}", "", f"观察：{_safe(atom['observation'])}", "", f"判分理由：{_safe(atom['reason'])}", ""]
                if atom["denominator"] is not None:
                    lines += [f"计算：{atom['numerator']} / {atom['denominator']} = {display(atom['state'])}", ""]
                    visible, omitted = _visible_items(atom["items"])
                    for item in visible:
                        if isinstance(item, dict):
                            label = _item_label(item) or item.get("source") or "核验项"
                            item_outcome = item.get("result") or item.get("verdict") or ""
                            location = item.get("location") or item.get("locator") or item.get("scope") or ""
                            note = item.get("reason") or item.get("explanation") or item.get("note") or item.get("basis") or ""
                            status = _item_passed(item)
                            verdict = "通过" if status is True else "未通过" if status is False else str(item_outcome)
                            detail = f"{label} — {verdict}"
                            if location:
                                detail += f"；位置：{location}"
                            if note:
                                detail += f"；依据：{note}"
                            if item.get("quote"):
                                detail += f"；摘录：{item['quote']}"
                            lines.append(f"- {_safe(detail)}")
                        else:
                            lines.append(f"- 核验项：{_safe(item)}")
                    if omitted:
                        lines.append(f"- 其余 {omitted} 条核验项见完整 JSON；此处未省略已识别的未通过项。")
                    lines.append("")
                display_notes = atom.get("validation_notes", [])
                if any(str(note).startswith("逐项清单共") for note in display_notes):
                    from .scoring import counting_notes
                    display_notes = counting_notes(atom.get("items", []), atom["denominator"])
                for note in display_notes:
                    lines.append(f"- 记录检查：{_safe(note)}")
                for correction in atom.get("record_corrections", []):
                    if isinstance(correction, dict):
                        detail = str(correction.get('reason', correction))
                        if 'original_denominator' in correction and 'audited_count' in correction:
                            detail += f" 原分母 {correction['original_denominator']}，实际计分样本 {correction['audited_count']}。"
                        if 'from' in correction and 'to' in correction:
                            detail += f" 判定由 {correction['from']} 改为 {correction['to']}。"
                    else:
                        detail = str(correction)
                    lines.append(f"- 记录修正：{_safe(detail)}")
                if display_notes:
                    lines.append("")
                for e in atom["evidence"]:
                    content = ("视觉观察：" + e["description"] if e.get("kind") == "visual" and e.get("description") else e.get("quote", ""))
                    lines.append(f"- 证据：`{_safe(e['file'])}` · {_safe(e['locator'])} · {_safe(content)}")
                lines.append("")
        if result["errors"]:
            lines += ["### 关键错误与封顶", ""]
            for error in result["errors"]:
                lines += [f"- {error['id']}：{'触发' if error['triggered'] else '未触发'}。{_safe(error['reason'])}"]
            lines.append("")
    return "\n".join(lines)


def summaries(tasks, result_map: dict, config: dict, run_dir: Path) -> dict:
    names = {s["id"]: s["name"] for s in config["subjects"]}
    process = process_records(config)
    by_dim = defaultdict(list)
    for task in tasks:
        results = [result_map[(task.id, s["id"])] for s in config["subjects"]]
        by_dim[task.dimension].append((task, results))
        path = run_dir / "results" / task.id / ("compare-result.md" if len(results) > 1 else "evaluation-result.md")
        write_text(path, task_markdown(task, results, names))
    dimensions = {}
    for dim, entries in by_dim.items():
        summary = {"dimension": dim, "tasks": [t.id for t, _ in entries], "subjects": {}, "task_results": [],
                   "model": config.get("model"), "aggregation": "任务等权，未舍入分数汇总，展示 ROUND_HALF_UP 两位小数"}
        for task, results in entries:
            row = {"task_id": task.id, "task_name": task.name, "subjects": {}}
            for result in results:
                losses = sorted((a for a in result["atoms"] if number(a["state"]) < 1), key=lambda a: (number(a["state"]), -number(a["weight"])))
                row["subjects"][result["subject"]] = {
                    "scores": result["scores"],
                    "main_losses": [{"atom_id": a["id"], "metric": a["metric"], "brief": _loss_brief(a)} for a in losses[:3]],
                    "critical_errors": result["errors"],
                    "verification_scope": [{"atom_id": a["id"], "numerator": a["numerator"],
                                             "denominator": a["denominator"],
                                             "scope_note": _short_reason(_summary_text(a["observation"]), 500)}
                                            for a in result["atoms"] if a["kind"] == "CLAIM-RATIO"],
                    "fingerprint": result["fingerprint"],
                    "evidence_record": {"path": f"{task.id}/{result['subject']}.json",
                                        "sha256": hashlib.sha256((run_dir / "results" / task.id / f"{result['subject']}.json").read_bytes()).hexdigest()},
                }
            baseline = config['subjects'][0]['id']
            row['differences'] = {s['id']: {m: str(number(row['subjects'][s['id']]['scores'][m]['final']) - number(row['subjects'][baseline]['scores'][m]['final'])) for m in METRICS} for s in config['subjects'][1:]}
            for s in config['subjects'][1:]:
                row['differences'][s['id']]['quality_mean'] = str(number(row['subjects'][s['id']]['scores']['quality_mean']) - number(row['subjects'][baseline]['scores']['quality_mean']))
            summary["task_results"].append(row)
        md = [f"# {_safe(dim)} 总评", "", f"纳入题目：{', '.join(summary['tasks'])}", "",
              f"评测模型：{_safe(summary['model'])}。统计：{summary['aggregation']}。", ""]
        md += ['## 逐题六项分数与差异', '', '差值为后列对象减首列对象；完成率差值单位为百分点，其余为分。质量均分只平均五项质量指标。', '',
               '| 题目 | 对象或差值 | ' + ' | '.join(LABELS[m] for m in METRICS) + ' | 质量均分 |',
               '| --- | --- | ' + ' | '.join('---:' for _ in range(7)) + ' |']
        for row in summary['task_results']:
            for s in config['subjects']:
                scores = row['subjects'][s['id']]['scores']
                values = [display(scores[m]['final']) + ('%' if m == 'completion' else '') for m in METRICS]
                md.append('| ' + row['task_id'] + ' | ' + _safe(s['name']) + ' | ' + ' | '.join(values + [display(scores['quality_mean'])]) + ' |')
            for s in config['subjects'][1:]:
                delta = row['differences'][s['id']]
                md.append('| ' + row['task_id'] + ' | ' + _safe(s['name'] + ' − ' + config['subjects'][0]['name']) + ' | ' + ' | '.join(display(delta[m]) for m in (*METRICS, 'quality_mean')) + ' |')
        md += ['']
        archived_process = {'record_type': 'subject_test_excel', 'dimension': dim, 'subjects': {}}
        for s in config["subjects"]:
            sid = s["id"]
            rows = [next(r for r in results if r["subject"] == sid) for _, results in entries]
            metric_scores = {}
            for metric in METRICS:
                values = [number(r["scores"][metric]["final"]) for r in rows]
                metric_scores[metric] = str(sum(values, Decimal(0)) / len(values))
            metric_scores["quality_mean"] = str(sum((number(metric_scores[m]) for m in METRICS[1:]), Decimal(0)) / 5)
            records = [process.get(t.id, {}).get(sid, {}) for t, _ in entries]
            durations = [x for record in records if (x := duration_seconds(record.get("runtime"))) is not None]
            successes = sum("成功" in str(r.get("status", "")) and not any(term in str(r.get("status", "")) for term in ("不成功", "未成功", "失败")) for r in records)
            process_stats = {"duration_count": len(durations), "mean_seconds": str(Decimal(sum(durations)) / len(durations)) if durations else None, "recorded_success_count": successes}
            summary["subjects"][sid] = {"name": s["name"], "scores": metric_scores}
            archived_process['subjects'][sid] = {'name': s['name'], 'task_ids': summary['tasks'], 'records': records, 'statistics': process_stats}
            md += [f"## {_safe(s['name'])}", "", f"任务数：{len(rows)}。任务完成率 {display(metric_scores['completion'])}%；质量均分 {display(metric_scores['quality_mean'])} 分。", "", "| 指标 | 均分 |", "| --- | ---: |"]
            md += [f"| {LABELS[m]} | {display(metric_scores[m])}{'%' if m == 'completion' else ''} |" for m in METRICS]
            md += ["", "### 逐题结果", "", "| 题目 | 完成率 | 质量均分 | 主要失分项 |", "| --- | ---: | ---: | --- |"]
            for row in summary["task_results"]:
                item = row["subjects"][sid]
                brief = item["main_losses"]
                md.append(f"| {row['task_id']} {_safe(row['task_name'])} | {display(item['scores']['completion']['final'])}% | {display(item['scores']['quality_mean'])} | {_safe('、'.join(x['atom_id'] for x in brief) or '无')} |")
            md += ["", "### 主要失分依据", ""]
            weak = sorted(((r["task_id"], a) for r in rows for a in r["atoms"] if number(a["state"]) < 1), key=lambda entry: (number(entry[1]["state"]), -number(entry[1]["weight"])))[:8]
            if weak:
                md += [f"- {task_id}/{a['id']}（{LABELS[a['metric']]}）：{_safe(_loss_brief(a))}" for task_id, a in weak]
            else:
                md.append("- 纳入题目中未发现原子项失分。")
            md += ["", "### 关键错误与封顶", ""]
            triggered = [(r, error) for r in rows for error in r["errors"] if error.get("triggered")]
            for result, error in triggered:
                caps = "、".join(f"{LABELS[m]}最终 {display(result['scores'][m]['final'])}" for m in METRICS if result["scores"][m]["caps"] or result["scores"][m].get("scoped_caps"))
                md.append(f"- {result['task_id']}/{error['id']}：{_safe(error['reason'])} 封顶后的指标：{caps}。")
            if not triggered:
                md.append("- 未触发关键错误封顶。")
            md += ["", "### 完整证据索引", ""]
            for row in summary["task_results"]:
                item = row["subjects"][sid]
                record = item["evidence_record"]
                md.append(f"- {row['task_id']}：[完整判分记录]({record['path']})；SHA-256 `{record['sha256']}`；输入指纹 `{item['fingerprint']}`。")
            md.append("")
        if len(config["subjects"]) > 1:
            base = config["subjects"][0]
            md += ["## 对象比较", "", f"差值方向：后列对象减 {_safe(base['name'])}。", "", "| 指标 | " + " | ".join(_safe(s["name"]) for s in config["subjects"]) + " |", "| --- | " + " | ".join("---:" for _ in config["subjects"]) + " |"]
            for metric in (*METRICS, "quality_mean"):
                md.append("| " + (LABELS.get(metric) or "质量均分") + " | " + " | ".join(display(summary["subjects"][s["id"]]["scores"][metric]) for s in config["subjects"]) + " |")
            md += ["", "相对首列的质量均分差值：" + "；".join(f"{s['name']} {display(number(summary['subjects'][s['id']]['scores']['quality_mean']) - number(summary['subjects'][base['id']]['scores']['quality_mean']))}" for s in config["subjects"][1:]) + "。", ""]
        slug = entries[0][0].id.split(".")[0]
        write_json(run_dir / 'process-records' / f'dimension-{slug}-subject-tests.json', archived_process)
        write_text(run_dir / "results" / f"dimension-{slug}-summary.md", "\n".join(md))
        write_json(run_dir / "results" / f"dimension-{slug}-summary.json", summary)
        dimensions[dim] = summary
    return dimensions


def make_report(config: dict, tasks, result_map: dict, dimensions: dict, run_dir: Path) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    for style in ("Title", "Heading 1", "Heading 2"):
        doc.styles[style].font.name = "Microsoft YaHei"
        doc.styles[style].font.color.rgb = __import__("docx").shared.RGBColor(0, 0, 0)
    title_style = doc.styles["Title"]
    title_style.font.size = Pt(21)
    title_style.font.bold = True
    for border in title_style.element.xpath("./w:pPr/w:pBdr"):
        border.getparent().remove(border)
    doc.add_paragraph(config.get("report_title", "AI Agent 应用能力评测报告"), style="Title")
    doc.add_paragraph(f"数据集：{Path(config['dataset']).name}  |  纳入 {len(tasks)} 道题  |  受测对象 {len(config['subjects'])} 个")
    doc.add_heading("评测范围与方法", level=1)
    doc.add_paragraph("本报告依据逐题 Rubric-Council 评分规则，对每个受测对象的交付独立判分。完成率与五项质量指标分别统计；各题等权，数据集统计直接平均题目原始分数。具体判分理由和证据见逐题结果文件。")
    doc.add_heading("总体结果", level=1)
    top_table = doc.add_table(rows=1, cols=3)
    top_table.style = "Table Grid"
    for i, label in enumerate(["受测对象", "任务完成率", "质量均分"]):
        top_table.rows[0].cells[i].text = label
    doc.add_paragraph("五项质量指标均分")
    metric_table = doc.add_table(rows=1, cols=6)
    metric_table.style = "Table Grid"
    for i, label in enumerate(["受测对象", "内容覆盖度", "准确率忠实度", "格式合规度", "结构完整度", "幻觉自洽性"]):
        metric_table.rows[0].cells[i].text = label
    for subject in config["subjects"]:
        sid = subject["id"]
        rows = [result_map[(t.id, sid)] for t in tasks]
        values = {m: sum((number(r["scores"][m]["final"]) for r in rows), Decimal(0)) / len(rows) for m in METRICS}
        mean = sum((values[m] for m in METRICS[1:]), Decimal(0)) / 5
        cells = top_table.add_row().cells
        for i, value in enumerate([subject["name"], display(values["completion"]) + "%", display(mean)]):
            cells[i].text = value
        cells = metric_table.add_row().cells
        for i, value in enumerate([subject["name"], *[display(values[m]) for m in METRICS[1:]]]):
            cells[i].text = value
    if len(config["subjects"]) > 1:
        base = config["subjects"][0]
        base_rows = [result_map[(t.id, base["id"])] for t in tasks]
        base_mean = sum((sum((number(r["scores"][m]["final"]) for r in base_rows), Decimal(0)) / len(base_rows) for m in METRICS[1:]), Decimal(0)) / 5
        for subject in config["subjects"][1:]:
            rows = [result_map[(t.id, subject["id"])] for t in tasks]
            mean = sum((sum((number(r["scores"][m]["final"]) for r in rows), Decimal(0)) / len(rows) for m in METRICS[1:]), Decimal(0)) / 5
            doc.add_paragraph(f"{subject['name']}相对{base['name']}的质量均分差值为 {display(mean - base_mean)} 分。")
    for dim, summary in dimensions.items():
        doc.add_heading(dim, level=1)
        doc.add_paragraph(f"纳入 {len(summary['tasks'])} 道题：{', '.join(summary['tasks'])}。")
        for subject in config["subjects"]:
            item = summary["subjects"][subject["id"]]
            score = item["scores"]
            doc.add_heading(subject["name"], level=2)
            doc.add_paragraph(f"任务完成率 {display(score['completion'])}%；质量均分 {display(score['quality_mean'])} 分。")
            weak = sorted(((t, a) for t in tasks if t.dimension == dim for a in result_map[(t.id, subject["id"])]["atoms"] if number(a["state"]) < 1), key=lambda pair: (number(pair[1]["state"]), -number(pair[1]["weight"])))[:3]
            if weak:
                doc.add_paragraph("主要失分依据：")
                for task, atom in weak:
                    paragraph = doc.add_paragraph(f"{task.id} {atom['id']}：{_loss_brief(atom)}", style="List Bullet")
                    paragraph.paragraph_format.keep_together = True
    if len(tasks) > 1:
        doc.add_heading("核验说明", level=1)
        doc.add_paragraph("报告中的分数由保存的逐原子状态按 rubric 权重计算。逐题 JSON 包含原始状态、分子分母、证据、理由和封顶判断；逐题 Markdown 便于人工复核。对于历史复用的交付，依据其执行时点判别时效性。")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = "Evaluation Judger 评测报告"
    target = run_dir / "reports" / "evaluation-report.docx"
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    return target
