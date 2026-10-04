"""Dataset discovery and test-process record intake."""
from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .rubric import Rubric, load_rubric


@dataclass(frozen=True)
class Subject:
    id: str
    name: str
    prefix: str


@dataclass(frozen=True)
class Task:
    id: str
    name: str
    dimension: str
    directory: Path
    question: Path
    rubric: Rubric
    sources: tuple[Path, ...]
    deliveries: dict[str, tuple[Path, ...]]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def load_config(path: Path) -> dict:
    config = json.loads(path.read_text(encoding="utf-8"))
    if not 1 <= len(config.get("subjects", [])) <= 3:
        raise ValueError("subjects must contain 1–3 entries")
    ids = [s["id"] for s in config["subjects"]]
    prefixes = [s["prefix"].lower() for s in config["subjects"]]
    if len(ids) != len(set(ids)) or len(prefixes) != len(set(prefixes)):
        raise ValueError("Subject IDs and prefixes must be unique")
    if not Path(config["dataset"]).is_dir():
        raise ValueError(f"Dataset does not exist: {config['dataset']}")
    return config


def discover(config: dict) -> list[Task]:
    root = Path(config["dataset"]).resolve()
    subjects = [Subject(**s) for s in config["subjects"]]
    excluded_prefixes = [p.lower() for p in config.get("excluded_subject_prefixes", [])]
    if any(p in [s.prefix.lower() for s in subjects] for p in excluded_prefixes):
        raise ValueError("Excluded subject prefix overlaps a selected subject")
    wanted = set(config.get("dimensions", []))
    layout = config.get("layout", {})
    dim_pattern = re.compile(layout.get("dimension_pattern", r"^维度(?P<id>[一二三四五六七八九十]+)-(?P<name>.+)$"))
    task_pattern = re.compile(layout.get("task_pattern", r"^(?P<id>\d+)_(?P<name>.+)$"))
    tasks: list[Task] = []
    for dim in sorted(p for p in root.iterdir() if p.is_dir()):
        dim_match = dim_pattern.match(dim.name)
        if not dim_match:
            continue
        dim_id = dim_match.group("id")
        digit = int(dim_id) if dim_id.isdigit() else "一二三四五六七八九十".index(dim_id) + 1 if len(dim_id) == 1 else None
        if digit is None:
            raise ValueError(f"Please configure an Arabic dimension ID for {dim.name}")
        if wanted and digit not in wanted:
            continue
        folders = [(p, task_pattern.match(p.name)) for p in dim.iterdir() if p.is_dir()]
        for folder, match in sorted((pair for pair in folders if pair[1]), key=lambda pair: int(pair[1].group("id"))):
            if not match:
                continue
            task_id = f"{digit}.{int(match.group('id'))}"
            if config.get("tasks") and task_id not in config["tasks"]:
                continue
            question = folder / layout.get("question_file", "问题描述.txt")
            rubric_path = folder / layout.get("rubric_file", "rubric.md")
            if not question.is_file() or not rubric_path.is_file():
                raise ValueError(f"Missing question or rubric: {folder}")
            deliveries: dict[str, list[Path]] = {s.id: [] for s in subjects}
            sources: list[Path] = []
            for p in sorted(x for x in folder.rglob("*") if x.is_file()):
                if p in (question, rubric_path) or p.name.startswith("~$"):
                    continue
                first_part = p.relative_to(folder).parts[0]
                found = [s for s in subjects if p.name.lower().startswith(s.prefix.lower()) or first_part.lower().startswith(s.prefix.lower())]
                if any(p.name.lower().startswith(prefix) or first_part.lower().startswith(prefix) for prefix in excluded_prefixes):
                    continue
                if len(found) > 1:
                    raise ValueError(f"Ambiguous participant: {p}")
                if found:
                    deliveries[found[0].id].append(p)
                else:
                    sources.append(p)
            if any(not deliveries[s.id] for s in subjects):
                missing = [s.id for s in subjects if not deliveries[s.id]]
                raise ValueError(f"Missing delivery for {task_id}: {missing}")
            tasks.append(Task(task_id, match.group("name"), dim.name, folder, question, load_rubric(rubric_path), tuple(sources), {k: tuple(v) for k, v in deliveries.items()}))
    if not tasks:
        raise ValueError("No tasks found in configured scope")
    dimension_order = {int(digit): index for index, digit in enumerate(config.get("dimensions", []))}
    def task_order(task):
        dimension, question = (int(part) for part in task.id.split("."))
        return dimension_order.get(dimension, dimension), question
    return sorted(tasks, key=task_order)


def process_records(config: dict) -> dict[str, dict[str, dict]]:
    """The layout is configurable; default matches DeepInsight's process sheet."""
    spec = config.get("process_excel", {})
    file = Path(spec.get("path", ""))
    if not file.is_file():
        raise ValueError("A process Excel file is required")
    wb = load_workbook(file, read_only=True, data_only=True)
    try:
        sheet = wb[spec.get("sheet", wb.sheetnames[0])]
        if sheet.max_row is None:
            # Some exporters omit OOXML dimension metadata. Read the actual
            # sheet bounds without changing the mandatory source workbook.
            sheet.calculate_dimension(force=True)
        task_col = spec.get("task_column", "C")
        columns = spec.get("subjects", {})
        task_id_map = spec.get("task_id_map", {})
        if not isinstance(task_id_map, dict):
            raise ValueError("process_excel.task_id_map must be an object")
        result: dict[str, dict[str, dict]] = {}
        for row in range(spec.get("first_row", 3), min(spec.get("last_row", sheet.max_row), sheet.max_row) + 1):
            raw = sheet[f"{task_col}{row}"].value
            if raw is None:
                continue
            task_label = str(raw).strip()
            task_id = str(task_id_map.get(task_label, task_label)).strip()
            if not re.fullmatch(r"\d+\.\d+", task_id):
                continue
            # Instruction/example sections may repeat a task ID without any
            # participant data. They must not overwrite a genuine test record.
            has_data = any(sheet[f"{col[key]}{row}"].value is not None
                           for col in columns.values() for key in ("runtime", "status", "executed_at")
                           if col.get(key))
            if not has_data:
                continue
            if task_id in result:
                raise ValueError(f"Duplicate process task {task_id} at {sheet.title}!{task_col}{row}")
            result[task_id] = {}
            for subject in config["subjects"]:
                col = columns.get(subject["id"], {})
                if not col:
                    raise ValueError(f"Missing process columns for {subject['id']}")
                result[task_id][subject["id"]] = {
                    "runtime": sheet[f"{col['runtime']}{row}"].value,
                    "status": sheet[f"{col['status']}{row}"].value,
                    "executed_at": sheet[f"{col['executed_at']}{row}"].value if col.get("executed_at") else config.get("execution_dates", {}).get(task_id, {}).get(subject["id"]),
                    "sheet": sheet.title,
                    "row": row,
                    "task_label": task_label,
                }
        return result
    finally:
        wb.close()


def collection_task_codes(task_id: str, config: dict) -> list[str]:
    aliases = config.get("process_excel", {}).get("task_id_map", {})
    return sorted({task_id, *(str(key) for key, value in aliases.items() if str(value) == task_id)})


def logical_delivery_name(name: str, task_id: str, subject_prefix: str, config: dict) -> str:
    """Remove confirmed collection labels only; preserve physical file names."""
    if not config.get("collection_prefixes_are_metadata", False) or not subject_prefix or not name.casefold().startswith(subject_prefix.casefold()):
        return name
    remainder = name[len(subject_prefix):]
    for code in collection_task_codes(task_id, config):
        if remainder.casefold().startswith((code + "-").casefold()):
            remainder = remainder[len(code)+1:]
            break
    return remainder or name


def task_fingerprint(task: Task, subject: str, config: dict) -> str:
    files = [task.question, task.rubric.path, *task.sources, *task.deliveries[subject]]
    payload = {str(p.relative_to(task.directory)): sha256(p) for p in files}
    payload["model"] = config.get("model", "aiaaa/deepseek-v4.1-flash#high")
    executed_at = config.get("execution_dates", {}).get(task.id, {}).get(subject)
    if executed_at is not None:
        payload["execution_date"] = str(executed_at)
    payload["judge_version"] = "0.1"
    if config.get("enable_image_input", False):
        payload["image_input_version"] = "1"
    if config.get("collection_prefixes_are_metadata", False):
        prefix = next(s["prefix"] for s in config["subjects"] if s["id"] == subject)
        payload["collection_naming"] = {"version": "1", "subject_prefix": prefix,
                                        "task_codes": collection_task_codes(task.id, config)}
    if "extraction_limit" in config:
        payload["extraction_limit"] = config["extraction_limit"]
    return hashlib.sha256(json.dumps(payload, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
